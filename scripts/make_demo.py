"""Build a self-contained Keystone demo: tracker data + an Obsidian vault full of realistic (fictional) engineering files.

    python scripts/make_demo.py                          # creates ./demo  (data/ + vault/ + meetings.ics)
    python scripts/make_demo.py --llm-from data/llm.json # also copy your AI model profiles into the demo
    python scripts/make_demo.py --force                  # rebuild it from scratch

Run Keystone on it without touching your real data:

    KEYSTONE_DATA_DIR=demo/data keystone-worklog --port 5050      # or: python app.py from a checkout

Dates are relative to today, so the demo always looks current. Needs the [documents] packages
(pypdf is not needed to *write* the PDFs; python-docx, python-pptx and openpyxl are).
"""
import argparse
import json
import random
import shutil
import subprocess
import sys
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from pathlib import Path

MARK = ".keystone-demo"
rng = random.Random(42)
TODAY = date.today()


def sunday_of(d):
    return d - timedelta(days=(d.weekday() + 1) % 7)


def iso(d):
    return d.isoformat()


# ---------- US federal holidays (as observed), for LH rows ----------
def federal_holidays(y):
    def nth(m, wd, n):  # wd: 0=Mon
        d = date(y, m, 1)
        d += timedelta(days=(wd - d.weekday()) % 7)
        return d + timedelta(weeks=n - 1)

    def last(m, wd):
        d = date(y, m + 1, 1) - timedelta(days=1) if m < 12 else date(y, 12, 31)
        return d - timedelta(days=(d.weekday() - wd) % 7)

    def obs(d):
        return d - timedelta(days=1) if d.weekday() == 5 else d + timedelta(days=1) if d.weekday() == 6 else d
    days = [obs(date(y, 1, 1)), nth(1, 0, 3), nth(2, 0, 3), last(5, 0), obs(date(y, 6, 19)), obs(date(y, 7, 4)), nth(9, 0, 1),
            nth(10, 0, 2), obs(date(y, 11, 11)), nth(11, 3, 4), obs(date(y, 12, 25))]
    return set(days)


HOLIDAYS = federal_holidays(TODAY.year - 1) | federal_holidays(TODAY.year) | federal_holidays(TODAY.year + 1)

# ---------- tracker data ----------
PROJECTS = [
    {"id": "bridge", "network": "4001234", "activity": "0010", "name": "Girder Bridge rehab", "budget": "400"},
    {"id": "runway", "network": "4005678", "activity": "0020", "name": "Runway lighting", "budget": "160"},
    {"id": "culvert", "network": "4009012", "activity": "0030", "name": "Culvert C-12 replacement", "budget": "100"},
    {"id": "training", "network": "4000001", "activity": "0090", "name": "Training", "budget": ""},
]


def row(p, code="RG", ts=False):
    return {"id": f"r{rng.randrange(10**8):08d}", "network": p["network"] if p else "", "activity": p["activity"] if p else "", "code": code, "ts": ts, "h": [""] * 7}


def make_period(start, last_day, leave_days=None):
    """Two Sun–Sat weeks; weekdays up to last_day get 8 h split across projects. Fridays are telework (TS)."""
    leave_days = leave_days or {}
    weeks = []
    for wi in range(2):
        rows = {}
        for di in range(7):
            d = start + timedelta(days=wi * 7 + di)
            if di in (0, 6) or d > last_day:
                continue
            if d in HOLIDAYS:
                rows.setdefault("LH", row(None, "LH"))["h"][di] = "8"
                continue
            if d in leave_days:
                code, hours = leave_days[d]
                rows.setdefault(code, row(None, code))["h"][di] = str(hours)
                if hours >= 8:
                    continue
            left = 8 - (leave_days.get(d, ("", 0))[1])
            ts = di == 5
            split = {"bridge": rng.choice([4, 4.5, 5, 5.5, 6]), "runway": rng.choice([0, 1, 1.5, 2]), "culvert": rng.choice([1, 1.5, 2])}
            if rng.random() < 0.15:
                split["training"] = 1
            total = sum(split.values())
            scale = left / total
            hours, acc = {}, 0
            for k, v in split.items():
                hours[k] = round(v * scale * 4) / 4
                acc += hours[k]
            hours["bridge"] += left - acc
            for pid, hv in hours.items():
                if hv <= 0:
                    continue
                p = next(x for x in PROJECTS if x["id"] == pid)
                key = pid + ("-ts" if ts else "")
                rows.setdefault(key, row(p, "RG", ts))["h"][di] = str(round(hv * 100) / 100)
        order = ["bridge", "bridge-ts", "runway", "runway-ts", "culvert", "culvert-ts", "training", "training-ts", "LA", "LS", "LH"]
        weeks.append({"rows": [rows[k] for k in order if k in rows] or [row(None)]})
    return {"weeks": weeks}


def tracker_data(vault):
    anchor = sunday_of(TODAY) - timedelta(days=7)          # the current pay period started last Sunday
    periods = {}
    for k in range(6, 0, -1):                                 # six finished pay periods…
        start = anchor - timedelta(days=14 * k)
        leave = {}
        if k == 4:
            leave[start + timedelta(days=5)] = ("LA", 8)        # a Friday of annual leave
        if k == 2:
            leave[start + timedelta(days=9)] = ("LS", 4)        # half a sick day
        per = make_period(start, start + timedelta(days=13), leave)
        if k > 1:                                             # …all submitted except the last one
            at = datetime.combine(start + timedelta(days=13), datetime.min.time()).replace(hour=16).isoformat() + "Z"
            total = sum(float(v or 0) for w in per["weeks"] for r in w["rows"] for v in r["h"])
            per["submitted"] = {"at": at, "total": total}
            per["lastSubmitted"] = {"at": at, "snapshot": json.dumps(per["weeks"])}
        periods[iso(start)] = per
    periods[iso(anchor)] = make_period(anchor, TODAY - timedelta(days=1))  # current period: filled up to yesterday

    def due(n):
        return iso(TODAY + timedelta(days=n))
    deliverables = [
        {"id": "d1", "title": "30% plans – Girder Bridge", "projectId": "bridge", "status": "done", "due": due(-40), "notes": ""},
        {"id": "d2", "title": "Bearing replacement details", "projectId": "bridge", "status": "doing", "due": due(3), "notes": "Pier 3 first; check jacking points."},
        {"id": "d3", "title": "Load rating memo", "projectId": "bridge", "status": "blocked", "due": due(-2), "notes": "Waiting on the girder FE run (girder_v3 didn't converge)."},
        {"id": "d4", "title": "60% plans – Girder Bridge", "projectId": "bridge", "status": "todo", "due": due(24), "notes": ""},
        {"id": "d5", "title": "Photometric calcs – runway edge lights", "projectId": "runway", "status": "doing", "due": due(6), "notes": ""},
        {"id": "d6", "title": "Lighting spec section 34 43 00", "projectId": "runway", "status": "todo", "due": due(15), "notes": ""},
        {"id": "d7", "title": "Culvert C-12 hydraulic memo", "projectId": "culvert", "status": "doing", "due": due(1), "notes": "100-year storm check; county wants it before storm season."},
        {"id": "d8", "title": "Culvert inspection report", "projectId": "culvert", "status": "done", "due": due(-12), "notes": ""},
    ]
    followups = [
        {"id": "f1", "from": "Maria Lopez (County)", "subject": "Culvert C-12 headwall", "received": due(-1), "due": due(1), "projectId": "culvert", "status": "open", "notes": "Confirm whether the 60 in pipe needs a new headwall.", "body": ""},
        {"id": "f2", "from": "Sam Ortiz (Fabricator)", "subject": "RE: bearing pads delivery", "received": due(-3), "due": due(-1), "projectId": "bridge", "status": "open", "notes": "", "body": ""},
        {"id": "f3", "from": "Airport ops", "subject": "Night closure window for lighting work", "received": due(-6), "due": due(4), "projectId": "runway", "status": "waiting", "notes": "", "body": ""},
        {"id": "f4", "from": "QA team", "subject": "Load rating checklist", "received": due(-20), "due": due(-15), "projectId": "bridge", "status": "done", "notes": "", "body": ""},
    ]
    away = [
        {"id": "a1", "kind": "travel", "start": due(8), "end": due(9), "note": "Site visit – Girder Bridge", "fill": False},
        {"id": "a2", "kind": "annual", "start": due(26), "end": due(28), "note": "Family trip", "fill": False},
    ]
    timer_log = [{"id": "t1", "date": due(-1), "network": "4009012", "activity": "0030", "code": "RG", "ts": False, "ms": int(1.7 * 36e5)}]
    settings = {"weekStart": "sun", "anchor": iso(anchor), "periodHours": 80, "leaveStart": iso(anchor - timedelta(days=14 * 6)),
                "sickRate": 4, "annualRate": 6, "sickStart": 52, "annualStart": 196, "vaultPath": str(vault.resolve()), "vaultFolder": "Keystone",
                "schedule": "5x8", "holidays": True, "carryLimit": 240, "captureNote": "Today.md", "dayFormat": "dddd, MMMM D, YYYY", "inboxFolder": "Inbox"}
    return {"settings": settings, "projects": PROJECTS, "deliverables": deliverables, "followups": followups, "away": away,
            "events": [], "timerLog": timer_log, "periods": periods, "timer": None}


# ---------- project plans (Plan tab) ----------
DAY_MIN = 480
PEOPLE = [{"uid": 1, "name": "Me", "initials": "ME", "me": True, "maxUnits": 1, "personId": "pe_jd01"},
          {"uid": 2, "name": "Maria Lopez", "initials": "ML", "maxUnits": 1, "personId": "pe_ml01"},
          {"uid": 3, "name": "Sam Ortiz", "initials": "SO", "maxUnits": 0.5, "away": [iso(TODAY + timedelta(days=12)), iso(TODAY + timedelta(days=13))], "personId": "pe_so01"},
          {"uid": 4, "name": "Priya Shah", "initials": "PS", "maxUnits": 1, "personId": "pe_ps01"}]
ME, MARIA, SAM, PRIYA = 1, 2, 3, 4


def plan_builder(nwa):
    """(tasks, assignments, add) – add(name, days, parent, preds, pct, people, **fields) appends a task and returns its uid."""
    tasks, assignments = [], []

    def add(name, days=None, parent=None, preds=(), pct=0, people=(), **kw):
        uid = len(tasks) + 1
        t = {"uid": uid, "name": name, "parentUid": parent, "links": [lnk(p) if isinstance(p, int) else p for p in preds],
             "percent": pct, "constraint": {"type": "ASAP"}, "baselines": {}, "notes": kw.pop("notes", ""), "priority": 500, "custom": {}}
        if days is not None:
            t["duration"] = {"minutes": int(days * DAY_MIN), "format": "d"}
            t["milestone"] = days == 0
            t["nwa"] = nwa
        t.update(kw)
        tasks.append(t)
        assignments.extend({"taskUid": uid, "resourceUid": who, "units": units} for who, units in people)
        return uid
    return tasks, assignments, add


def lnk(pred, kind="FS", lag_days=0):
    return {"pred": pred, "type": kind, "lag": {"minutes": int(lag_days * DAY_MIN), "format": "d"}}


def plan_shell(name, project_id, start, tasks, assignments, raid=(), resources=PEOPLE, **kw):
    return {"version": 1, "name": name, "projectId": project_id, "start": iso(start) + "T08:00", "statusDate": None,
            "units": {"minutesPerDay": 480, "minutesPerWeek": 2400, "daysPerMonth": 20}, "calendar": None, "calendarSource": "keystone",
            "tasks": tasks, "resources": [dict(r) for r in resources], "assignments": assignments, "changeLog": [], "raid": list(raid), **kw}


def demo_plan(baseline=False):
    """The Girder Bridge rehab plan: phases, milestones, links, people, progress and a few issues and risks.
    baseline=True gives the original estimates (shorter analysis and design) – saved as baseline 0 so the slip shows."""
    start = sunday_of(TODAY) - timedelta(days=34)              # a Monday five weeks ago
    tasks, assignments, add = plan_builder("4001234|0010")
    p1 = add("Inspection and data")
    a = add("Site inspection", 3, p1, pct=100, people=[(ME, 1), (MARIA, 1)])
    b = add("Survey and as-builts", 4 if baseline else 5, p1, [a], pct=100, people=[(SAM, 1)])
    c = add("Inspection report", 4, p1, [b], pct=100, people=[(ME, 1)])
    m1 = add("Inspection complete", 0, p1, [c], pct=100)
    p2 = add("Analysis")
    d = add("Girder FE model (Abaqus)", 6 if baseline else 8, p2, [m1], pct=100, people=[(ME, 1)], notes="girder_v3 didn't converge – v4 with finer mesh at the bearings.")
    e = add("Load rating", 6, p2, [d], pct=60, people=[(ME, 1), (PRIYA, 0.5)], budget=11000)
    f = add("Bearing replacement design", 5 if baseline else 7, p2, [lnk(m1, "FS", 2)], pct=85, people=[(MARIA, 1)], budget=9500)
    g = add("Load rating memo", 2, p2, [e], pct=0, people=[(ME, 1)], deadline=iso(TODAY + timedelta(days=12)) + "T17:00")
    p3 = add("Design")
    add("30% plans submitted", 0, p3, [m1], pct=100)
    h_ = add("60% plans", 12, p3, [f, lnk(e, "SS", 3)], pct=0 if baseline else 30, people=[(MARIA, 1), (PRIYA, 1)], budget=32000)
    i = add("Specifications", 6, p3, [lnk(h_, "FF")], people=[(PRIYA, 1)], budget=6500)
    j = add("Cost estimate", 4, p3, [h_], people=[(SAM, 1)])
    m3 = add("60% plans submitted", 0, p3, [h_, i, j])
    k = add("County review", 15, p3, [m3], people=[], notes="The county takes about three weeks.")
    l_ = add("90% plans", 10, p3, [k], people=[(MARIA, 1), (PRIYA, 1)])
    add("PS&E submitted", 0, p3, [l_, g], deadline=iso(TODAY + timedelta(days=95)) + "T17:00")
    return plan_shell("Girder Bridge rehab", "bridge", start, tasks, assignments, raid=[
                {"id": "r1", "kind": "risk", "title": "FE model may not converge in time for the load rating", "owner": "Me", "due": iso(TODAY + timedelta(days=5)), "status": "open",
                 "probability": 3, "impact": 4, "mitigation": "Fall back to a line-girder rating for the memo.", "taskUids": [d, e], "created": iso(TODAY - timedelta(days=6)), "description": ""},
                {"id": "r2", "kind": "issue", "title": "Bearing pads delivery slipped two weeks", "owner": "Sam Ortiz", "due": iso(TODAY + timedelta(days=2)), "status": "open",
                 "probability": 0, "impact": 3, "mitigation": "", "taskUids": [f], "created": iso(TODAY - timedelta(days=3)), "description": "Fabricator says mid-month at the earliest."},
                {"id": "r3", "kind": "risk", "title": "County review takes longer than three weeks", "owner": "Maria Lopez", "due": "", "status": "monitoring",
                 "probability": 2, "impact": 3, "mitigation": "Pre-submittal meeting with the county.", "taskUids": [k], "created": iso(TODAY - timedelta(days=20)), "description": ""},
                {"id": "r4", "kind": "decision", "title": "Replace bearings at Pier 3 first", "owner": "Me", "due": "", "status": "closed", "closed": iso(TODAY - timedelta(days=10)),
                 "probability": 0, "impact": 0, "mitigation": "", "taskUids": [f], "created": iso(TODAY - timedelta(days=12)), "description": "",
                 "decisionStatus": "decided", "decided": iso(TODAY - timedelta(days=10)), "decidedBy": "Maria Lopez (County)",
                 "options": ["Pier 3 first, then piers 1–2", "All three piers in one closure"], "rationale": "Pier 3 bearings are the worst (frozen, 40 mm out of position); doing them first keeps the other piers open to traffic and fits the May–June closure window.",
                 "decisionImpact": {"days": 0, "cost": None, "scope": ""}},
                {"id": "r5", "kind": "decision", "title": "Use a line-girder rating for the memo if the FE model isn't ready", "owner": "Maria Lopez", "due": iso(TODAY + timedelta(days=3)), "status": "open",
                 "probability": 0, "impact": 0, "mitigation": "", "taskUids": [g], "created": iso(TODAY - timedelta(days=2)), "description": "Keeps the memo deadline; the FE rating follows in the 60% submittal.",
                 "decisionStatus": "proposed", "options": ["Wait for the FE model", "Line-girder rating now, FE later"]},
                {"id": "r6", "kind": "risk", "title": "Fabricator may claim extra for expediting the bearing pads", "owner": "Me", "due": iso(TODAY + timedelta(days=9)), "status": "open",
                 "probability": 3, "impact": 2, "mitigation": "Hold to the contract dates; document the delay cause.", "taskUids": [f], "created": iso(TODAY - timedelta(days=3)), "description": "", "internal": True}])


def runway_plan(baseline=False):
    """Runway edge lighting: a smaller plan that is behind – the circuit survey is late and the closure window isn't confirmed."""
    start = sunday_of(TODAY) - timedelta(days=13)              # a Monday two weeks ago
    tasks, assignments, add = plan_builder("4005678|0020")
    p1 = add("Existing conditions")
    a = add("Record drawings review", 2, p1, pct=100, people=[(ME, 1)])
    b = add("Existing circuit survey", 4 if baseline else 6, p1, [a], pct=50 if not baseline else 0, people=[(SAM, 1)],
            notes="Night work only – two nights lost to weather.")
    p2 = add("Design")
    c = add("Photometric calcs", 8, p2, [a], pct=70, people=[(ME, 1)])
    d = add("Lighting layout", 5, p2, [b, c], people=[(PRIYA, 1)])
    e = add("Spec section 34 43 00", 4, p2, [lnk(d, "SS", 2)], people=[(ME, 0.5)], deadline=iso(TODAY + timedelta(days=15)) + "T17:00")
    m1 = add("Design complete", 0, p2, [d, e])
    p3 = add("Coordination")
    f = add("Confirm night closure window", 3, p3, [a], pct=30, people=[(MARIA, 0.5)])
    g = add("FAA review", 10, p3, [m1, f], notes="Airport ops forward it to the FAA district office.")
    add("Bid package issued", 0, p3, [g])
    return plan_shell("Runway lighting", "runway", start, tasks, assignments, raid=[
        {"id": "r1", "kind": "issue", "title": "Night closure window not confirmed by airport ops", "owner": "Maria Lopez", "due": iso(TODAY + timedelta(days=4)),
         "status": "open", "probability": 0, "impact": 4, "mitigation": "", "taskUids": [f, g], "created": iso(TODAY - timedelta(days=6)), "description": ""},
        {"id": "r2", "kind": "risk", "title": "Existing cable is direct-buried and needs replacing", "owner": "Sam Ortiz", "due": iso(TODAY + timedelta(days=7)),
         "status": "open", "probability": 2, "impact": 5, "mitigation": "Megger test during the survey; carry an add-alternate in the bid.", "taskUids": [b, d],
         "created": iso(TODAY - timedelta(days=9)), "description": ""}])


def culvert_plan():
    """Culvert C-12 replacement – written as an MS Project XML file to try File → Open an MS Project file."""
    start = sunday_of(TODAY) + timedelta(days=8)               # a Monday next week
    tasks, assignments, add = plan_builder("4009012|0030")
    p1 = add("Hydraulics")
    a = add("100-year storm check", 3, p1, people=[(ME, 1)])
    b = add("Hydraulic memo", 2, p1, [a], people=[(ME, 1)])
    p2 = add("Design")
    c = add("Pipe and headwall design", 5, p2, [b], people=[(MARIA, 1)])
    d = add("Erosion control plan", 3, p2, [lnk(c, "SS", 2)], people=[(PRIYA, 1)])
    e = add("Plans and quantities", 4, p2, [c, d], people=[(MARIA, 1), (PRIYA, 0.5)])
    add("Submit to the county", 0, p2, [e])
    return plan_shell("Culvert C-12 replacement", "culvert", start, tasks, assignments,
                      resources=[r for r in PEOPLE if r["uid"] != SAM])


def bridge_template():
    """A reusable plan for the next bridge rehab: phases, tasks and links but no dates, progress or people (New plan → Template)."""
    tasks, assignments, add = plan_builder("")
    p1 = add("Inspection and data")
    a = add("Site inspection", 3, p1)
    b = add("Survey and as-builts", 5, p1, [a])
    c = add("Inspection report", 4, p1, [b])
    m1 = add("Inspection complete", 0, p1, [c])
    p2 = add("Analysis")
    d = add("FE model", 8, p2, [m1])
    e = add("Load rating", 6, p2, [d])
    f = add("Load rating memo", 2, p2, [e])
    p3 = add("Design")
    g = add("60% plans", 12, p3, [lnk(e, "SS", 3)])
    h_ = add("Specifications", 6, p3, [lnk(g, "FF")])
    m2 = add("60% plans submitted", 0, p3, [g, h_])
    k = add("Agency review", 15, p3, [m2])
    l_ = add("90% plans", 10, p3, [k])
    add("PS&E submitted", 0, p3, [l_, f])
    for t in tasks:
        t.pop("nwa", None)
    return plan_shell("Bridge rehab template", None, sunday_of(TODAY) + timedelta(days=8), tasks, [], resources=PEOPLE[:1], template=True)


# The scheduler lives in the page (static/planner/engine.js). When Node is installed it fills in the dates and
# baselines here too, so the demo opens with grey baseline bars and the MS Project file carries real dates.
STAMP_JS = r"""
import { readFileSync } from 'node:fs';
const E = await import(process.argv[1]);
const C = await import(process.argv[2]);
const { plans, holidays, dir, model } = JSON.parse(readFileSync(0, 'utf8'));
const y = new Date().getFullYear();
const cal = E.calendarFromSchedule([0, 8, 8, 8, 8, 8, 0, 0, 8, 8, 8, 8, 8, 0], '2026-01-04', Object.fromEntries(holidays.map(d => [d, 'Holiday'])), `${y - 1}-01-01`, `${y + 2}-12-31`);
const stamp = p => { const s = E.schedule({ ...p, calendar: cal }); for (const t of p.tasks) { const r = s.byUid.get(t.uid); t.start = E.fromMin(r.es); t.finish = E.fromMin(r.ef); } return s; };
for (const { plan, baseline } of plans) {
  stamp(plan);
  if (baseline) {
    const bs = E.schedule({ ...baseline, calendar: cal }), b = E.setBaseline({ ...baseline, calendar: cal }, bs, 0);
    C.baselineCosts(b, bs, C.planContext(b, dir, model), 0);  // planned cost at the original estimates
    const byUid = new Map(b.tasks.map(t => [t.uid, t.baselines]));
    for (const t of plan.tasks) t.baselines = byUid.get(t.uid) || {};
  }
}
process.stdout.write(JSON.stringify(plans.map(p => p.plan)));
"""


def stamp_dates(pairs):
    """[(plan, baseline plan or None)] → the plans with start/finish (and baseline 0) filled in, or unchanged without Node."""
    node = shutil.which("node")
    engine = Path(__file__).resolve().parent.parent / "keystone_worklog" / "static" / "planner" / "engine.js"
    if not node or not engine.exists():
        return [p for p, _ in pairs], False
    doc = people_doc()
    payload = json.dumps({"plans": [{"plan": p, "baseline": b} for p, b in pairs], "holidays": sorted(iso(d) for d in HOLIDAYS), "dir": doc, "model": doc["costModels"][0]})
    r = subprocess.run([node, "--input-type=module", "-e", STAMP_JS, engine.as_uri(), (engine.parent / "costs.js").as_uri()], input=payload, capture_output=True, text=True)
    if r.returncode:
        print("Couldn't schedule the demo plans with Node (they still work, without baselines):", r.stderr.strip()[-300:])
        return [p for p, _ in pairs], False
    return json.loads(r.stdout), True



# ---------- team costs: People & rates, ERP exports, a cost ledger, status history ----------
STAFF = [  # id, directory name, how the ERP writes it, employee ID, category, pay per hour, plan resource uid
    ("pe_jd01", "Jordan Doe", "DOE, JORDAN", "00398812", "Engineer IV", 58.40, ME),
    ("pe_ml01", "Maria Lopez", "LOPEZ, MARIA", "00412873", "Senior Engineer", 65.00, MARIA),
    ("pe_so01", "Sam Ortiz", "ORTIZ, SAMUEL", "00377120", "Survey Technician", 36.50, SAM),
    ("pe_ps01", "Priya Shah", "SHAH, PRIYA", "00451009", "Engineer II", 44.00, PRIYA),
]
FRINGE = 0.31
LAYOUT_HEAD = ["Posting Date", "Work Date", "Employee Name", "Employee ID", "Network", "Activity", "Hours", "Pay Type", "Pay Amount", "Fringe Amount", "Cost Element", "Document No.", "Description"]


def people_doc():
    """data/people.json: the team, dated pay rates, two cost models and the master ERP layout."""
    jan = iso(date(TODAY.year, 1, 4))
    people = [{"id": i, "name": n, "aliases": [erp], "employeeId": emp, "category": cat, "maxUnits": 0.5 if i == "pe_so01" else 1,
               "away": [iso(TODAY + timedelta(days=12)), iso(TODAY + timedelta(days=13))] if i == "pe_so01" else [], "me": i == "pe_jd01", "active": True}
              for i, n, erp, emp, cat, _, _ in STAFF]
    rates = [{"who": i, "base": r if i != "pe_ml01" else 62.50, "from": jan} for i, _, _, _, _, r, _ in STAFF]
    rates.append({"who": "pe_ml01", "base": 65.00, "from": iso(date(TODAY.year, 7, 5))})  # Maria's July raise
    rates += [{"who": "cat:CAD Technician", "base": 38.00, "from": jan}, {"who": "default", "base": 45.00, "from": jan}]
    charges = [{"name": "Overhead", "kind": "pctBase", "value": 85}, {"name": "Service charge", "kind": "pctRunning", "value": 6}, {"name": "IT & admin", "kind": "perHour", "value": 4.25}]
    return {"rev": 1, "people": people, "rates": rates,
            "costModels": [{"id": "cm_county", "name": "County contract", "base": "erp", "fringePct": FRINGE * 100, "adders": charges, "payTypes": {"OT": 1.5}, "nonLaborPct": 5, "isDefault": True},
                           {"id": "cm_rate", "name": "Hours × rate (estimates)", "base": "rate", "fringePct": FRINGE * 100, "adders": charges, "payTypes": {"OT": 1.5}, "nonLaborPct": 0}],
            "layouts": [{"id": "ly_demo01", "name": "Labor distribution report", "master": True, "dateOrder": "mdy", "decimalComma": False,
                         "fields": {"post": ["Posting Date"], "date": ["Work Date"], "person": ["Employee Name"], "emp": ["Employee ID"], "nwa": ["Network", "Activity"],
                                    "hours": ["Hours"], "payType": ["Pay Type"], "pay": ["Pay Amount"], "fringe": ["Fringe Amount"], "category": ["Cost Element"],
                                    "doc": ["Document No."], "desc": ["Description"]}}],
            "settings": {"currency": "$", "dateOrder": "mdy"}}


def posting_date(d):
    """The ERP posts a week's labor on the Saturday that ends it."""
    return d + timedelta(days=(5 - d.weekday()) % 7)


def erp_lines(plans, periods):
    """Charge lines for everyone: other people from the plans' schedules, you from your own timesheet."""
    rate = {i: r for i, _, _, _, _, r, _ in STAFF}
    who = {res: (i, erp, emp) for i, _, erp, emp, _, _, res in STAFF}
    lines, doc = [], 4900012000
    overrun = {MARIA: 1.15, SAM: 1.1, PRIYA: 1.1}

    def pay_for(pid, d, hrs, ptype="RG"):
        r = rate[pid] if not (pid == "pe_ml01" and d < date(TODAY.year, 7, 5)) else 62.50
        pay = round(hrs * r * (1.5 if ptype == "OT" else 1), 2)
        return pay, round(pay * FRINGE, 2)

    for plan, nwa in plans:
        for a in plan["assignments"]:
            if a["resourceUid"] == ME:
                continue
            t = next(x for x in plan["tasks"] if x["uid"] == a["taskUid"])
            if not t.get("start") or not t.get("duration") or t.get("milestone"):
                continue
            d0, d1 = date.fromisoformat(t["start"][:10]), date.fromisoformat(t["finish"][:10])
            done_by = d1 if t["percent"] >= 100 else min(d1, TODAY - timedelta(days=1)) if t["percent"] > 0 or d0 < TODAY else d0 - timedelta(days=1)
            d = d0
            while d <= done_by:
                if d.weekday() < 5 and d not in HOLIDAYS:
                    pid, erp, emp = who[a["resourceUid"]]
                    hrs = round(8 * a["units"] * overrun.get(a["resourceUid"], 1) * 4) / 4
                    pay, fr = pay_for(pid, d, hrs)
                    doc += 1
                    lines.append({"date": d, "name": erp, "emp": emp, "nwa": nwa, "hours": hrs, "payType": "RG", "pay": pay, "fringe": fr, "category": "Labor", "doc": str(doc), "desc": t["name"]})
                d += timedelta(days=1)
    nwas = {f"{p['network']}|{p['activity']}" for p in PROJECTS}
    for start, per in periods.items():  # your hours: straight from the timesheet
        for wi, w in enumerate(per["weeks"]):
            for r in w["rows"]:
                key = f"{r['network']}|{r['activity']}"
                if r["code"] != "RG" or key not in nwas or key == "4000001|0090":
                    continue
                for di, v in enumerate(r["h"]):
                    if not v or float(v) <= 0:
                        continue
                    d = date.fromisoformat(start) + timedelta(days=wi * 7 + di)
                    pay, fr = pay_for("pe_jd01", d, float(v))
                    doc += 1
                    lines.append({"date": d, "name": "DOE, JORDAN", "emp": "00398812", "nwa": key, "hours": float(v), "payType": "RG", "pay": pay, "fringe": fr, "category": "Labor", "doc": str(doc), "desc": "Timesheet"})
    for l in lines:
        l["post"] = posting_date(l["date"])
    return lines, doc


def erp_rows(lines):
    rows = []
    for l in sorted(lines, key=lambda x: (x["post"], x["nwa"], x["name"] or "", x["date"])):
        net, act = l["nwa"].split("|")
        rows.append([l["post"].strftime("%m/%d/%Y"), l["date"].strftime("%m/%d/%Y"), l["name"], l["emp"], net, act, l["hours"], l["payType"],
                     l["pay"], l["fringe"], l["category"], l["doc"], l["desc"]])
    return rows


def erp_exports(dest, data, plans, periods):
    """Two ERP exports (an older one already imported, and a newer one that overlaps it), the ledgers, and budgets."""
    lines, doc = erp_lines(plans, periods)
    sat = TODAY - timedelta(days=(TODAY.weekday() - 5) % 7 or 7)   # the last Saturday before today: posted through
    cut_a = sat - timedelta(days=21)
    a = [l for l in lines if l["post"] <= cut_a]
    b = [l for l in lines if cut_a - timedelta(days=13) <= l["post"] <= sat]
    # the newer export corrects the older one: a reversal and a re-post, plus overtime, a travel bill and someone new
    fix = next((l for l in b if l["name"] == "ORTIZ, SAMUEL" and l["post"] <= cut_a), None)
    if fix:
        doc += 1
        b.append({**fix, "hours": -fix["hours"], "pay": -fix["pay"], "fringe": -fix["fringe"], "doc": str(doc), "desc": "Reversal – wrong network", "post": cut_a + timedelta(days=7)})
        doc += 1
        b.append({**fix, "nwa": "4005678|0020", "doc": str(doc), "desc": "Re-posted to runway lighting", "post": cut_a + timedelta(days=7)})
    late = sat - timedelta(days=9)
    doc += 1
    b.append({"date": late, "post": sat, "name": "SHAH, PRIYA", "emp": "00451009", "nwa": "4001234|0010", "hours": 3, "payType": "OT", "pay": round(3 * 44 * 1.5, 2), "fringe": round(3 * 44 * 1.5 * FRINGE, 2), "category": "Labor", "doc": str(doc), "desc": "60% plans – evening push"})
    for k in range(3):
        d = sat - timedelta(days=12 - k)
        if d.weekday() < 5:
            doc += 1
            b.append({"date": d, "post": sat, "name": "NGUYEN, LINH", "emp": "00466230", "nwa": "4001234|0010", "hours": 6, "payType": "RG", "pay": 228.0, "fringe": 70.68, "category": "Labor", "doc": str(doc), "desc": "CAD support – 60% plans"})
    doc += 1
    b.append({"date": sat - timedelta(days=15), "post": sat, "name": "", "emp": "", "nwa": "4001234|0010", "hours": "", "payType": "", "pay": 412.80, "fringe": "", "category": "Travel", "doc": str(doc), "desc": "Site visit mileage and lodging"})
    title = lambda through: [["Labor Distribution – Project Charges"], [f"Posting dates through {through.strftime('%m/%d/%Y')} · run {TODAY.strftime('%m/%d/%Y')}"], []]
    name_a, name_b = f"ERP labor report {cut_a.strftime('%b %d')}.xlsx", f"ERP labor report {sat.strftime('%b %d')}.xlsx"
    make_xlsx(dest / name_a, [("Charges", title(cut_a) + [LAYOUT_HEAD] + erp_rows(a))])
    make_xlsx(dest / name_b, [("Charges", title(sat) + [LAYOUT_HEAD] + erp_rows(b))])
    # the older export is already in the ledgers (as if imported last month)
    costs = data / "costs"
    costs.mkdir()
    for p in PROJECTS:
        key = f"{p['network']}|{p['activity']}"
        mine = [l for l in a if l["nwa"] == key]
        if not mine:
            continue
        canon = [{k: v for k, v in {"date": iso(l["date"]), "post": iso(l["post"]), "name": l["name"], "emp": l["emp"], "hours": l["hours"], "pay": l["pay"], "fringe": l["fringe"],
                                    "payType": l["payType"], "category": l["category"], "nwa": key, "doc": l["doc"], "desc": l["desc"], "batch": "b_demo0001"}.items() if v not in ("", None)} for l in mine]
        led = {"projectId": p["id"], "postedThrough": iso(cut_a), "rev": 1, "lines": canon,
               "batches": [{"id": "b_demo0001", "source": name_a, "layout": "Labor distribution report", "imported": iso(cut_a + timedelta(days=3)) + "T09:12:00",
                            "covers": {"from": min(x["post"] for x in canon), "to": iso(cut_a)}, "rows": len(canon), "replaced": 0,
                            "hours": round(sum(x.get("hours", 0) for x in canon), 2), "pay": round(sum(x.get("pay", 0) for x in canon), 2)}]}
        f = costs / f"{p['id']}.json"
        f.write_text(json.dumps(led), encoding="utf-8")
        f.chmod(0o600)
    return name_a, name_b, cut_a, sat


def status_history(plan, numbers):
    """Three weekly status snapshots for the bridge plan, with the milestones slipping a little each week."""
    snaps = []
    friday = TODAY - timedelta(days=(TODAY.weekday() - 4) % 7 or 7)
    leafs = [t for t in plan["tasks"] if t.get("duration")]
    for k, frac in enumerate([0.3, 0.6, 0.85]):
        when = friday - timedelta(days=7 * (2 - k))
        shift = lambda t, key="finish": iso(date.fromisoformat(t["baselines"]["0"][key][:10]) + timedelta(days=round((date.fromisoformat(t[key][:10]) - date.fromisoformat(t["baselines"]["0"][key][:10])).days * frac)))
        ms = {str(t["uid"]): shift(t) for t in leafs if t.get("milestone") and t.get("baselines", {}).get("0")}
        pct_then = lambda t: 100 if t["percent"] >= 100 and t["finish"][:10] <= iso(when) else min(t["percent"], round(t["percent"] * (0.55 + 0.2 * k) / 5) * 5)
        tasks = {str(t["uid"]): [t["start"][:10], shift(t) if t.get("baselines", {}).get("0") else t["finish"][:10], pct_then(t), t["name"]] for t in leafs}
        bac = numbers["bac"]
        pct = round(numbers["pct"] * (0.62 + 0.15 * k)); exp = pct + 5 + 3 * k
        pv, ev = bac * exp / 100, bac * pct / 100
        cpi = [0.97, 0.93, 0.9][k]
        ac = ev / cpi
        rag = [("green", "On plan."), ("amber", "FE model running late; load rating start moved."), ("amber", "FE model 2 days late; 60% submittal pushed.")][k]
        snaps.append({"id": "s_" + iso(when), "date": iso(when), "statusDate": iso(when),
                      "rag": {"overall": rag[0], "schedule": rag[0], "cost": "green" if cpi > 0.95 else "amber", "scope": "green",
                              "reasons": {"overall": rag[1], "schedule": rag[1] if rag[0] != "green" else "", "cost": "" if cpi > 0.95 else f"CPI {cpi:.2f} (below 0.95)"}},
                      "finish": max(ms.values()) if ms else plan["tasks"][-1]["finish"][:10], "baselineFinish": numbers["baselineFinish"], "pct": pct, "pctExpected": exp,
                      "milestones": ms, "tasks": tasks, "postedThrough": iso(when - timedelta(days=6)),
                      "ev": {"bac": round(bac), "pv": round(pv), "ev": round(ev), "ac": round(ac), "spi": round(ev / pv, 2), "cpi": cpi, "eac": round(ac + (bac - ev) / cpi),
                             "evH": None, "acH": None, "budget": numbers["budget"], "runOut": None},
                      "raid": [{"id": x["id"], "kind": x["kind"], "title": x["title"], "score": x["probability"] * x["impact"], "owner": x["owner"], "due": x["due"]} for x in plan["raid"] if x["status"] != "closed" and (k > 0 or x["id"] != "r2")],
                      "openRisks": 2, "openIssues": 0 if k == 0 else 1,
                      "done": [["Inspection report", "◆ Inspection complete"], ["◆ 30% plans submitted"], ["Bearing pads delivery confirmed"]][k],
                      "next": [["Girder FE model (Abaqus)", "Bearing replacement design"], ["Girder FE model – finer mesh at the bearings", "Load rating"], ["Load rating", "60% plans"]][k],
                      "decisions": [] if k < 2 else ["Approve the line-girder fallback for the rating memo"], "notes": "", "by": "you"})
    return snaps


# ---------- the project brief (Plan → Brief) ----------
def demo_brief():
    return """---
keystone: brief
plan: p_0000b01d
summary: Rehabilitate the three-span steel girder bridge on County Road 12 – new bearings at pier 3, deck repair and a load rating.
client: County of Example – Public Works
cover: ""
---
# Girder Bridge rehab – project brief

The living summary of this project. Edit it here or in Keystone (Plan → Brief). The numbers – status, schedule, costs, risks – are added by Keystone when you share the brief.

## Background

The three-span steel girder bridge on County Road 12 over Mill Creek (built 1968) failed its last inspection at pier 3: the elastomeric bearings are frozen and the girders have walked about 40 mm. The county wants the bearings replaced, the deck joints repaired and a current load rating before the 2027 construction season.

## Objectives

- Replace the bearings at pier 3, then piers 1 and 2
- Repair the deck joints and the spalled deck areas
- Issue a load rating memo for the posted and legal loads
- Deliver 60%, 90% and final PS&E on the county's schedule

## Scope

### In scope
- Inspection, survey and as-builts
- Girder FE model and load rating
- Bearing replacement and jacking design
- Deck repair details, specifications and cost estimate

### Out of scope
- Painting and lead abatement (separate county contract)
- Approach roadway and guardrail
- Construction inspection

## Success criteria

- PS&E accepted by the county by the contract date
- Load rating memo accepted without a re-submittal
- Design within the $156,000 fee

## Assumptions and constraints

- Night and weekend closures only, May to June
- One lane open at all times during jacking
- County review takes about three weeks per submittal

## Stakeholders

| Who | Cares about | How often they hear from us |
| --- | --- | --- |
| County Public Works | Schedule, closures, cost | Weekly status, monthly meeting |
| Mill Creek residents | Closures and detours | Before each closure |
| Bearing fabricator | Shop drawings, delivery dates | As needed |

## Key facts

- **Client:** County of Example – Public Works
- **Contract:** 26-114
- **NWA:** 4001234 / 0010
- **Location:** County Road 12 over Mill Creek
- **Funding:** County bond + state bridge program
- **Structure:** Three-span continuous steel girder, 1968

## Key links

- [[Projects/Girder Bridge rehab/Girder Bridge rehab.md|Project notes]]
- [[Projects/Girder Bridge rehab/Specs/Girder Bridge special provisions.pdf|Special provisions]]
- [[Projects/Girder Bridge rehab/Budget tracker.xlsx|Budget tracker]]
- [[Analysis/README.md|Analysis notes]]

## Contacts

| Name | Role | Organisation | Email | Phone |
| --- | --- | --- | --- | --- |
| Maria Lopez | County project manager | County of Example | mlopez@example.gov | 555-0142 |
| Sam Ortiz | Bearings | Ortiz Steel | sam@ortiz-steel.example | 555-0187 |
| Priya Shah | Design engineer | Our office | pshah@example.com | |

## Lessons learned

- Ask the fabricator for pad lead times at 30%, not 60%
"""


def demo_log():
    """A few weeks of updates for the bridge plan's weekly log."""
    def at(days, hh=10):
        return iso(TODAY - timedelta(days=days)) + f"T{hh:02d}:00"
    items = [
        (24, "contact", "Kick-off with the county: closures only May–June, one lane open during jacking.", False),
        (20, "site", "Site visit with Maria: pier 3 bearings frozen, girders walked ~40 mm. Photos in the project folder.", False),
        (16, "win", "Inspection report submitted a day early.", False),
        (12, "decision", "County agreed to do pier 3 first – see the decision log.", False),
        (9, "problem", "girder_v3 FE run didn't converge at the jacking step – rebuilding the mesh at the bearings (v4).", False),
        (6, "update", "FE model v4 converges; load rating started.", False),
        (4, "problem", "Bearing pads delivery slipped two weeks at the fabricator.", False),
        (3, "update", "Staffing: Priya moves to 60% plans full-time from next week; Linh helps with CAD.", True),
        (2, "win", "County approved the pier 3 jacking detail.", False),
        (1, "contact", "Maria asked for a line-girder rating fallback so the memo date holds.", False),
    ]
    return [{"id": f"l_demo{i:02d}", "at": at(d, 9 + i % 6), "kind": k, "text": t, "internal": internal, "by": "you"} for i, (d, k, t, internal) in enumerate(items)]

# ---------- documents ----------
def make_pdf(pages):
    """A small PDF with a real text layer (what an OCR'd scan looks like to a reader)."""
    def esc(t):  # the built-in Helvetica only has Latin-1
        t = t.replace("\u2013", "-").replace("\u2014", "-").replace("\u2019", "'").encode("latin-1", "replace").decode("latin-1")
        return t.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    objs = ["<< /Type /Catalog /Pages 2 0 R >>",
            f"<< /Type /Pages /Kids [{' '.join(f'{4 + 2 * i} 0 R' for i in range(len(pages)))}] /Count {len(pages)} >>",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, text in enumerate(pages):
        lines = text.split("\n")
        stream = "BT /F1 16 Tf 72 730 Td 22 TL (" + esc(lines[0]) + ") Tj T* /F1 11 Tf 15 TL " + " ".join(f"({esc(l)}) Tj T*" for l in lines[1:]) + " ET"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>")
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode() + b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    return out + f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()


def make_docx(path, title, sections, table=None):
    import docx
    d = docx.Document()
    d.add_heading(title, level=0)
    for head, paras in sections:
        d.add_heading(head, level=1)
        for p in paras:
            d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        t.style = "Table Grid"
        for i, r in enumerate(table):
            for j, v in enumerate(r):
                t.cell(i, j).text = str(v)
    d.save(path)


def make_pptx(path, slides):
    import pptx
    prs = pptx.Presentation()
    for title, body, notes in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    prs.save(path)


def make_xlsx(path, sheets):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)


def make_eml(path, frm, to, subject, when, body, attachments=()):
    m = EmailMessage()
    m["From"], m["To"], m["Subject"], m["Date"] = frm, to, subject, when.strftime("%a, %d %b %Y %H:%M:%S -0700")
    m.set_content(body)
    for name, data, sub in attachments:
        m.add_attachment(data, maintype="application", subtype=sub, filename=name)
    path.write_bytes(m.as_bytes())


def outlook_txt(frm, when, subject, body):
    return f"From: {frm}\nSent: {when.strftime('%A, %B %d, %Y %I:%M %p')}\nTo: Hiren\nSubject: {subject}\n\n{body}\n"


ABAQUS_INP = """*Heading
** Girder Bridge rehab – span 2 girder model, {ver}
** Elastomeric bearing pads at pier 3 modelled as springs; jacking load case for bearing replacement
*Preprint, echo=NO, model=NO, history=NO
*Part, name=GIRDER
*Node
{nodes}*Element, type=C3D8R
*Material, name=Steel_A709_50
*Elastic
200000., 0.3
*Plastic
345., 0.
450., 0.15
*Material, name=Elastomer
*Hyperelastic, neo hooke
0.9, 0.002
*Step, name=Jacking, nlgeom=YES, inc=200
*Static
0.05, 1., 1e-08, 0.1
*Cload
JACK_PT, 2, -850000.
*End Step
"""


def abaqus_msg(converged, ver):
    lines = [f" Abaqus 2023 – job girder_{ver}", "", " STEP 1  INCREMENT 1  STEP TIME 0.000E+00", "       CONVERGENCE CHECKS FOR EQUILIBRIUM ITERATION   1"]
    lines += [" ***WARNING: 12 NODES ARE MISSING DEGREE OF FREEDOM CONSTRAINTS. THEY HAVE BEEN IDENTIFIED IN NODE SET WarnNodeMissMassOrInertia."] * 3
    if converged:
        lines += ["", " ITERATION SUMMARY FOR THE INCREMENT: 4 TOTAL ITERATIONS", "", " THE ANALYSIS HAS COMPLETED SUCCESSFULLY"]
    else:
        lines += [" ***WARNING: SOLVER PROBLEM. NUMERICAL SINGULARITY WHEN PROCESSING NODE GIRDER.1043 D.O.F. 2 RATIO = 1.2E+15"] * 4
        lines += [" ***ERROR: TOO MANY ATTEMPTS MADE FOR THIS INCREMENT", " ***ERROR: TIME INCREMENT REQUIRED IS LESS THAN THE MINIMUM SPECIFIED", "", " THE ANALYSIS HAS NOT BEEN COMPLETED"]
    return "\n".join(lines) + "\n"


def abaqus_sta(converged):
    rows = [" SUMMARY OF JOB INFORMATION:", " STEP  INC ATT SEVERE EQUIL TOTAL  TOTAL      STEP       INC OF", "              DISCON ITERS ITERS  TIME/    TIME/LPF    TIME/LPF"]
    t = 0.0
    for inc in range(1, 9 if converged else 5):
        dt = 0.125 if converged else 0.05 / inc
        t += dt
        rows.append(f"   1 {inc:5d}   1     0     {rng.randint(2, 5)}     {rng.randint(2, 5)}  {t:.4f}     {t:.4f}     {dt:.4E}")
    rows.append(" THE ANALYSIS HAS COMPLETED SUCCESSFULLY" if converged else " THE ANALYSIS HAS NOT BEEN COMPLETED")
    return "\n".join(rows) + "\n"


def abaqus_dat(ver):
    head = f" ABAQUS DATA FILE – girder_{ver}\n\n   P R O B L E M   S I Z E\n\n NUMBER OF ELEMENTS IS     48210\n NUMBER OF NODES IS        55902\n\n"
    table = "".join(f"  {i:6d}   {rng.uniform(-3, 3):.4E}  {rng.uniform(-250, 250):.4E}  {rng.uniform(0, 0.002):.4E}\n" for i in range(1, 6000))
    warn = " ***WARNING: THE STRAIN INCREMENT HAS EXCEEDED FIFTY TIMES THE STRAIN TO CAUSE FIRST YIELD AT 3 POINTS\n"
    return head + "   NODE FOOT-   U2           S22          PEEQ\n" + table + warn + table + " THE ANALYSIS HAS COMPLETED SUCCESSFULLY\n"


def write(vault, rel, content):
    p = vault / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(content, bytes):
        p.write_bytes(content)
    else:
        p.write_text(content, encoding="utf-8")
    return p


def build_vault(v):
    d = lambda n: TODAY + timedelta(days=n)  # noqa: E731
    fmt = lambda x: x.strftime("%Y-%m-%d")  # noqa: E731
    write(v, ".obsidian/templates.json", json.dumps({"folder": "Templates"}))
    write(v, "Templates/Meeting.md", "---\ntype: meeting\ndate: {{date}}\n---\n# {{title}}\n\n**Attendees:** \n\n## Notes\n\n## Action items\n- [ ] \n")
    write(v, "Templates/Daily.md", "# {{date}}\n\n## Focus\n\n## Notes\n")
    write(v, "Home.md", "# Home\nWork vault for bridge, airfield and drainage projects.\n\n- [[Girder Bridge rehab]]\n- [[Runway Lighting]]\n- [[Culvert C-12]]\n- Analysis models live in `Analysis/`, one folder per model.\n- New files go in `Inbox/` and get filed from there.\n")

    # --- Girder Bridge rehab ---
    gb = "Projects/Girder Bridge rehab"
    write(v, f"{gb}/Girder Bridge rehab.md", "# Girder Bridge rehab\nRehabilitation of the three-span steel girder bridge on County Road 12: bearing replacement at pier 3, deck repair and a load rating.\n\nNetwork 4001234 / activity 0010. County contract 26-114.\n\n## Contacts\n- Maria Lopez – county PM\n- Sam Ortiz – bearing fabricator\n\n## Status\n- [x] 30% plans\n- [ ] Bearing replacement details\n- [ ] Load rating memo\n")
    write(v, f"{gb}/Bearing replacement.md", "# Bearing replacement\nReplace the elastomeric bearing pads at pier 3. Jack the girders no more than 1/8 inch, one girder line at a time.\n\n## Open questions\n- Jacking points: stiffener check needed at each girder end.\n- Pad thickness 2 in laminated, 60 durometer.\n")
    write(v, f"{gb}/Deck repair.md", "# Deck repair\nRemove unsound concrete to sound material, patch with Class AA concrete. Estimate 340 sq ft of partial-depth repair on span 2.\n")
    write(v, f"{gb}/Specs/Girder Bridge special provisions.pdf", make_pdf([
        "SPECIAL PROVISIONS\nGirder Bridge Rehabilitation, County Road 12\nContract 26-114\nIssued for bid",
        "Section 412 – Concrete Deck Repair\nRemove unsound concrete to sound material by chipping hammer no heavier than 35 lb.\nPatch with Class AA concrete; wet cure for 7 days.",
        "Section 564 – Bearing Replacement\nReplace the elastomeric bearing pads at pier 3 with laminated pads, 60 durometer.\nJack the girders no more than 1/8 inch. Submit a jacking procedure 14 days before work.",
        "Section 620 – Traffic Control\nNight lane closures only, between 9 PM and 5 AM.\nOne lane open at all times; flaggers required during jacking."]))
    make_docx(v / f"{gb}/Scope of work.docx", "Scope of work – Girder Bridge rehab", [
        ("Scope", ["Replace the bearings at pier 3 of the Girder Bridge.", "Repair the deck on span 2 and prepare a load rating."]),
        ("Schedule", ["30% plans complete. 60% plans due next month.", "Night closures between May and June."])],
        [["Item", "Hours", "Cost"], ["Bearing replacement design", "120", "$18,000"], ["Deck repair", "80", "$12,000"], ["Load rating", "90", "$13,500"]])
    make_xlsx(v / f"{gb}/Budget tracker.xlsx", [("Budget", [["Task", "Hours", "Owner", "Status"], ["Bearing details", 120, "Hiren", "In progress"], ["Deck repair", 80, "Maria", "Not started"],
                                                            ["Load rating", 90, "Hiren", "Blocked"], ["Scour analysis", 64, "Priya", "Done"]]),
                                                 ("Notes", [["Contingency held at 10 percent"], ["Rates per 2026 agreement"]])])

    # --- Runway lighting ---
    rw = "Projects/Runway Lighting"
    write(v, f"{rw}/Runway Lighting.md", "# Runway Lighting\nReplace the runway 9-27 edge lights with LED fixtures and new constant-current regulators.\n\nNetwork 4005678 / activity 0020.\n\n- Night closures need airport ops approval.\n- Photometrics per FAA AC 150/5340-30.\n")
    write(v, f"{rw}/Lighting spec notes.md", "# Lighting spec notes\nSection 34 43 00: L-861 elevated edge lights, 200 ft spacing, CCR sized for 6.6 A.\n")
    make_pptx(v / f"{rw}/Runway lighting kickoff.pptx", [
        ("Runway 9-27 edge lighting", "LED retrofit of 112 elevated edge lights\nNew CCRs in the airfield lighting vault", "Mention the night closure windows."),
        ("Schedule", "Design: this quarter\nConstruction: next summer, night work only", ""),
        ("Risks", "Closure windows\nLong lead time on CCRs", "Ask ops for the NOTAM process.")])

    # --- Culvert C-12 ---
    cv = "Projects/Culvert C-12"
    write(v, f"{cv}/README.md", "# Culvert C-12\nReplacement of culvert C-12 under County Road 12: hydraulics, inspection reports and county correspondence.\n")
    write(v, f"{cv}/Culvert C-12.md", "# Culvert C-12\nThe existing 48 inch CMP is undersized for the 100-year storm. Proposed: 60 inch RCP with new headwalls.\n\nNetwork 4009012 / activity 0030.\n")
    write(v, f"{cv}/Hydraulics/Culvert C-12 hydraulic calcs.pdf", make_pdf([
        "Culvert C-12 Hydraulic Calculations\nDesign storm: 100-year, 24-hour\nPeak flow 212 cfs",
        "Existing 48 in CMP: headwater 8.9 ft, overtops the road at 180 cfs.\nProposed 60 in RCP: headwater 6.1 ft at 212 cfs, no overtopping."]))

    # --- meetings, mail, analysis, scripts ---
    meetings = [(d(-34), "County kickoff", "Kickoff with the county for the Girder Bridge and culvert C-12. County wants the culvert memo before storm season."),
                (d(-21), "Bearing pad coordination", "Sam confirmed laminated pads, 60 durometer. Lead time 6 weeks."),
                (d(-14), "Runway lighting design review", "Ops prefers closures Tue–Thu nights. CCR sizing to be checked."),
                (d(-7), "Load rating check-in", "Girder FE model v3 did not converge at the jacking step – numerical singularity at node 1043."),
                (d(-3), "County meeting", "Discussed the culvert headwall and the bearing jacking procedure. Action: send the hydraulic memo.")]
    for when, title, text in meetings:
        write(v, f"Meetings/{fmt(when)} {title}.md", f"---\ntype: meeting\ndate: {fmt(when)}\n---\n# {title}\n\n## Notes\n{text}\n\n## Action items\n- [ ] Follow up\n")
    write(v, "Mail/Mail.md", "# Mail\nSaved emails (.eml or text) about any project, one file per email.\n")
    make_eml(v / "Mail/Culvert sizing.eml", "Maria Lopez <maria.lopez@county.example>", "hiren@example.com", "Culvert C-12 sizing",
             datetime.combine(d(-9), datetime.min.time()).replace(hour=9, minute=15),
             "Hi,\nThe county wants the culvert memo by Friday, ahead of storm season.\nThanks, Maria",
             [("C-12 design memo.pdf", make_pdf(["Culvert C-12 design memo\nThe 48 inch pipe is undersized for the 100-year storm."]), "pdf")])
    write(v, "Mail/RE bearing pads delivery.txt", outlook_txt("Sam Ortiz", datetime.combine(d(-3), datetime.min.time()).replace(hour=16, minute=10), "RE: bearing pads delivery",
                                                             "The fabricator says the pads ship next week, so the jacking crew can start on the 16th."))
    af = "Analysis/Girder FE"
    write(v, "Analysis/README.md", "# Analysis\nFinite element models, one folder per model. Keep each Abaqus job's .inp, .msg, .sta, .dat and .odb together.\n")
    nodes = "".join(f"{i}, {i * 0.25:.2f}, 0., 0.\n" for i in range(1, 400))
    for ver, ok in (("v2", True), ("v3", False)):
        write(v, f"{af}/girder_{ver}.inp", ABAQUS_INP.format(ver=ver, nodes=nodes))
        write(v, f"{af}/girder_{ver}.msg", abaqus_msg(ok, ver))
        write(v, f"{af}/girder_{ver}.sta", abaqus_sta(ok))
        write(v, f"{af}/girder_{ver}.odb", b"\x00\x01ODB binary results\x00" * 500)
    write(v, f"{af}/girder_v2.dat", abaqus_dat("v2"))
    write(v, "Scripts/loads.py", '"""Load cases for the girder model."""\nimport math\n\n\ndef dead_load(span, w):\n    """Midspan moment of a simply supported span (kip-ft)."""\n    return w * span ** 2 / 8\n\n\ndef jacking_force(reaction, factor=1.3):\n    """Jacking force with the 30 % allowance from Section 564."""\n    return reaction * factor\n\n\nclass LoadCase:\n    def __init__(self, name, factor):\n        self.name, self.factor = name, factor\n')
    write(v, "Scripts/scour.ipynb", json.dumps({"nbformat": 4, "nbformat_minor": 5, "metadata": {}, "cells": [
        {"cell_type": "markdown", "metadata": {}, "source": ["# Scour depth at pier 3\n", "HEC-18 contraction scour, 100-year flow."]},
        {"cell_type": "code", "metadata": {}, "execution_count": 1, "source": ["y1, y0 = 6.2, 4.8\n", "ys = y1 - y0\n", "print(f'contraction scour {ys:.1f} ft')"],
         "outputs": [{"output_type": "stream", "name": "stdout", "text": ["contraction scour 1.4 ft\n"]}]}]}, indent=1))

    # --- today's captures ---
    heading = TODAY.strftime("%A, %B %-d, %Y") if sys.platform != "win32" else TODAY.strftime("%A, %B %#d, %Y")
    write(v, "Today.md", f"# Today\n\nQuick captures from Keystone, one section per day (newest first).\n\n## {heading}\n- 09:10 Call Maria about the culvert headwall #culvert\n- [ ] 11:30 Check jacking points on the girders @task due:fri\n")

    # --- the inbox: files to be filed ---
    ib = "Inbox"
    write(v, f"{ib}/girder_v4.inp", ABAQUS_INP.format(ver="v4 – refined mesh at the pier 3 bearings", nodes=nodes))
    write(v, f"{ib}/girder_v4.msg", abaqus_msg(True, "v4"))
    write(v, f"{ib}/girder_v4.sta", abaqus_sta(True))
    write(v, f"{ib}/RE culvert headwall.txt", outlook_txt("Maria Lopez", datetime.combine(d(-1), datetime.min.time()).replace(hour=15, minute=2), "RE: culvert headwall",
                                                          "The county asks whether the C-12 culvert upgrade to a 60 inch pipe also needs a new headwall before the storm season."))
    write(v, f"{ib}/bearing shop drawing review.txt", "Shop drawing review – laminated elastomeric bearing pads for pier 3, 2 in thick, 60 durometer.\nCheck the jacking points on the girders and the anchor bolt layout. Returned: approved as noted.\n")
    write(v, f"{ib}/Culvert C-12 inspection photos log.pdf", make_pdf([
        "Culvert C-12 Inspection Log\nInlet: headwall cracked, 2 ft scour hole at the apron.\nBarrel: 48 in CMP, invert corroded through at 3 locations.\nOutlet: riprap displaced."]))
    make_docx(v / f"{ib}/Pier 3 RFI response.docx", "RFI 07 response – pier 3 bearings", [
        ("Question", ["Contractor asks whether the girders may be jacked 1/4 inch at pier 3."]),
        ("Response", ["No. Jack no more than 1/8 inch per Section 564 of the special provisions. Submit the jacking procedure."])])
    write(v, f"{ib}/runway edge light photometrics.csv", "Fixture,Station,Intensity_cd,Required_cd,Pass\n" + "".join(
        f"L-861-{i:03d},{i * 200},{rng.randint(55, 75)},50,yes\n" for i in range(1, 30)))
    write(v, f"{ib}/grocery list.txt", "milk\neggs\nbread\ncoffee\n")


def meetings_ics(path):
    def dt(day, hh, mm=0):
        return f"{(TODAY + timedelta(days=day)).strftime('%Y%m%d')}T{hh:02d}{mm:02d}00"
    start = sunday_of(TODAY) - timedelta(days=14)
    events = [
        ("standup", "Daily stand-up", f"{start.strftime('%Y%m%d')}T090000", f"{start.strftime('%Y%m%d')}T091500", "RRULE:FREQ=WEEKLY;BYDAY=MO,WE,FR", "Teams"),
        ("review", "Girder Bridge design review", dt(0, 14), dt(0, 15, 30), "", "Room 4"),
        ("county", "County coordination call – culvert C-12", dt(1, 11), dt(1, 12), "", "Phone"),
        ("runway", "Runway lighting submittal review", dt(3, 10), dt(3, 11, 30), "", "Airport ops"),
        ("training", "Abaqus contact modelling training", dt(5, 13), dt(5, 16), "", "Online"),
    ]
    lines = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Keystone demo//EN"]
    for uid, title, s, e, rrule, loc in events:
        lines += ["BEGIN:VEVENT", f"UID:{uid}@demo", f"SUMMARY:{title}", f"DTSTART:{s}", f"DTEND:{e}", f"LOCATION:{loc}"] + ([rrule] if rrule else []) + ["END:VEVENT"]
    lines.append("END:VCALENDAR")
    path.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dest", type=Path, default=Path(__file__).resolve().parent.parent / "demo", help="output folder (default: ./demo)")
    ap.add_argument("--force", action="store_true", help="replace an existing demo folder")
    ap.add_argument("--llm-from", type=Path, help="copy AI model profiles from this llm.json (e.g. data/llm.json)")
    ap.add_argument("--no-git", action="store_true", help="don't make the demo vault a git repository")
    args = ap.parse_args()
    dest = args.dest.resolve()
    if dest.exists():
        if not (dest / MARK).exists():
            sys.exit(f"{dest} exists and isn't a Keystone demo folder – refusing to touch it.")
        if not args.force:
            sys.exit(f"{dest} already exists. Use --force to rebuild it.")
        shutil.rmtree(dest)
    data, vault = dest / "data", dest / "vault"
    data.mkdir(parents=True)
    vault.mkdir()
    (dest / MARK).write_text("Created by scripts/make_demo.py – safe to delete.\n")
    try:
        build_vault(vault)
    except ImportError as e:
        sys.exit(f"Missing package ({e.name}). Install the document packages first: pip install python-docx python-pptx openpyxl")
    (data / "plans").mkdir()
    plans, stamped = stamp_dates([(demo_plan(), demo_plan(baseline=True)), (runway_plan(), runway_plan(baseline=True)), (culvert_plan(), None), (bridge_template(), None)])
    bac = lambda p: sum((t.get("baselines", {}).get("0") or {}).get("cost") or 0 for t in p["tasks"])
    budgets = {"bridge": round(bac(plans[0]) * 1.06, -3) or 120000, "runway": round(bac(plans[1]) * 1.1, -3) or 60000, "culvert": 18000}
    for proj in PROJECTS:
        if proj["id"] in budgets:
            proj.update(budgetAmount=str(int(budgets[proj["id"]])), costModelId="cm_county")
    tracker = tracker_data(vault)
    (data / "data.json").write_text(json.dumps(tracker, indent=2), encoding="utf-8")
    meetings_ics(dest / "meetings.ics")
    now = datetime.now().strftime("%Y-%m-%dT%H:%M:%S")
    for pid, plan in zip(["p_0000b01d", "p_00009a1f", None, "p_0000c3e1"], plans):
        if pid:
            plan.update(id=pid, rev=1, updated=now)
            (data / "plans" / f"{pid}.json").write_text(json.dumps(plan, indent=1), encoding="utf-8")
    people = data / "people.json"
    people.write_text(json.dumps(people_doc(), indent=1), encoding="utf-8")
    people.chmod(0o600)
    (data / "plans" / "p_0000b01d.log.json").write_text(json.dumps({"entries": demo_log()}), encoding="utf-8")
    write(vault, "Keystone/Projects/Girder Bridge rehab/Brief.md", demo_brief())
    erp_a, erp_b, cut_a, posted = erp_exports(dest, data, [(plans[0], "4001234|0010"), (plans[1], "4005678|0020")], tracker["periods"])
    for proj, plan in (("bridge", plans[0]), ("runway", plans[1])):  # budgets also cover what was charged before the plan started
        led = json.loads((data / "costs" / f"{proj}.json").read_text()) if (data / "costs" / f"{proj}.json").exists() else {"lines": []}
        early = [l for l in led["lines"] if l["date"] < plan["start"][:10]]
        full = sum((l.get("pay", 0) + l.get("fringe", 0)) * 1.85 * 1.06 + l.get("hours", 0) * 4.25 for l in early)
        p = next(x for x in tracker["projects"] if x["id"] == proj)
        p["budgetAmount"] = str(int(round(int(p["budgetAmount"]) + full * 1.05, -3)))
    (data / "data.json").write_text(json.dumps(tracker, indent=2), encoding="utf-8")
    if stamped:
        b = plans[0]
        leafs = [t for t in b["tasks"] if t.get("duration") and not t.get("milestone")]
        tot = sum(t["duration"]["minutes"] for t in leafs) or 1
        nums = {"bac": bac(b), "budget": budgets["bridge"], "pct": round(sum(t["duration"]["minutes"] * t["percent"] for t in leafs) / tot),
                "baselineFinish": max(t["baselines"]["0"]["finish"][:10] for t in b["tasks"] if t.get("baselines", {}).get("0"))}
        st = data / "plans" / "p_0000b01d.status.json"
        st.write_text(json.dumps({"snapshots": status_history(b, nums)}), encoding="utf-8")
        st.chmod(0o600)
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from keystone_worklog.msproject import to_mspdi
    (dest / "Culvert C-12 schedule.xml").write_text(to_mspdi(plans[2]), encoding="utf-8")
    if args.llm_from:
        s = json.loads(args.llm_from.read_text(encoding="utf-8"))
        s["semantic"], s["documents"] = True, True
        f = data / "llm.json"
        f.write_text(json.dumps(s, indent=2), encoding="utf-8")
        f.chmod(0o600)
    if not args.no_git and shutil.which("git"):
        (vault / ".gitignore").write_text(".trash/\n.obsidian/workspace*.json\n*.odb\n", encoding="utf-8")
        g = ["git", "-c", "user.name=Keystone demo", "-c", "user.email=demo@localhost"]
        subprocess.run(["git", "init", "-q"], cwd=vault, check=True)
        subprocess.run([*g, "add", "-A"], cwd=vault, check=True)
        subprocess.run([*g, "commit", "-qm", "Demo vault"], cwd=vault, check=True)
    n = sum(1 for p in vault.rglob("*") if p.is_file() and ".git" not in p.parts)
    print(f"Demo ready in {dest}\n  data/   tracker data{' + AI profiles' if args.llm_from else ''}\n  vault/  {n} files ({len(list((vault / 'Inbox').iterdir()))} in Inbox/)\n  meetings.ics  to try Board → Calendar → Import calendar\n  data/plans/   2 project plans + a template (Plan tab){'' if stamped else ' – install Node.js for baselines'}\n  Culvert C-12 schedule.xml  to try Plan → Open an MS Project file\n  {erp_b}  to try Plan → Costs → Import from ERP (overlaps {erp_a}, already imported)\n\n"
          f"Run it (your real data is not touched):\n  KEYSTONE_DATA_DIR={data} keystone-worklog --port 5050\n")


if __name__ == "__main__":
    main()
