"""Team costs from the ERP: reading an export (Excel, CSV or pasted rows) and the per-project cost ledger.

Reading only turns the file into rows of cells. Choosing columns, parsing values, matching people and
working out the full cost all happen in the page (static/planner/costs.js), so the user sees every step.

A ledger is one JSON file per project (data/costs/<projectId>.json). Each import is a batch. A new batch
replaces the lines of earlier batches whose *posting* date falls in the range it covers, because an ERP
export is the truth for the period it was run for (corrections and reversals included). The lines it
replaced are kept in data/costs/.trash/, so the most recent import can be undone exactly.
"""
import base64
import csv
import datetime as dt
import io
import json
import re
import secrets
import time
from pathlib import Path

from . import fsutil

try:
    import openpyxl
except ImportError:  # pragma: no cover - openpyxl is a dependency, but keep CSV and paste working without it
    openpyxl = None

MAX_ROWS = 100_000
MAX_COLS = 80
PROJECT_ID = re.compile(r"^[A-Za-z0-9_-]{1,64}$")
ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")
PRIVATE = 0o600  # pay data: readable by this user only


class CostError(Exception):
    pass


# ---------- reading an export ----------
def _cell(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return v if v == v and abs(v) != float("inf") else ""  # no NaN/inf in JSON
    return str(v).strip()


def _trim(rows):
    """Drop trailing empty cells and rows; cap the size."""
    out = []
    for r in rows[:MAX_ROWS]:
        r = [_cell(c) for c in list(r)[:MAX_COLS]]
        while r and r[-1] == "":
            r.pop()
        out.append(r)
    while out and not out[-1]:
        out.pop()
    return out


def _decode(data):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _delimited(text):
    sample = text[:20000]
    if "\t" in sample:
        delim = "\t"
    else:
        try:
            delim = csv.Sniffer().sniff(sample, delimiters=",;|\t").delimiter
        except csv.Error:
            delim = ","
    return list(csv.reader(io.StringIO(text), delimiter=delim))


def read_table(filename="", data_b64=None, text=None, sheet=None):
    """An export → {"sheets": [names], "sheet": name, "rows": [[cell, …], …]}. Cells are text, numbers or ISO dates."""
    if text is not None:
        return {"sheets": [], "sheet": None, "rows": _trim(_delimited(text))}
    try:
        data = base64.b64decode(data_b64 or "")
    except ValueError:
        raise CostError("The file couldn't be read.") from None
    if not data:
        raise CostError("The file is empty.")
    suffix = Path(filename or "").suffix.lower()
    if suffix in (".xlsx", ".xlsm") or data[:2] == b"PK":
        if openpyxl is None:
            raise CostError("Reading Excel files needs openpyxl: pip install openpyxl (or save the sheet as CSV).")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise CostError(f"This isn't a readable Excel file ({e}).") from None
        try:
            names = [ws.title for ws in wb.worksheets]
            ws = wb[sheet] if sheet in names else None
            if ws is None:  # the first sheet with something in it
                for w in wb.worksheets:
                    if any(any(c not in (None, "") for c in r) for r in w.iter_rows(max_row=50, values_only=True)):
                        ws = w
                        break
                ws = ws or wb.worksheets[0]
            rows = _trim(list(ws.iter_rows(values_only=True, max_row=MAX_ROWS)))
            return {"sheets": names, "sheet": ws.title, "rows": rows}
        finally:
            wb.close()
    if suffix == ".xls" or data[:4] == b"\xd0\xcf\x11\xe0":
        raise CostError("Old .xls files can't be read – save it as .xlsx or CSV in Excel first.")
    return {"sheets": [], "sheet": None, "rows": _trim(_delimited(_decode(data)))}


# ---------- the ledger ----------
class Ledgers:
    def __init__(self, folder):
        self.dir = Path(folder)

    def path(self, pid):
        if not PROJECT_ID.match(pid or ""):
            raise CostError("Bad project id")
        return self.dir / f"{pid}.json"

    def load(self, pid):
        led = fsutil.read_json(self.path(pid), default=None)
        if not isinstance(led, dict):
            led = {}
        led.setdefault("projectId", pid)
        led.setdefault("postedThrough", None)
        led.setdefault("batches", [])
        led.setdefault("lines", [])
        led.setdefault("rev", 0)
        return led

    def _save(self, led):
        led["rev"] = (led.get("rev") or 0) + 1
        fsutil.write_text(self.path(led["projectId"]), json.dumps(led), mode=PRIVATE)

    def summary(self):
        out = []
        for f in sorted(self.dir.glob("*.json")) if self.dir.is_dir() else []:
            if not PROJECT_ID.match(f.stem):
                continue
            led = fsutil.read_json(f, default={}) or {}
            lines = led.get("lines") or []
            last = (led.get("batches") or [None])[-1]
            out.append({"projectId": f.stem, "postedThrough": led.get("postedThrough"), "lines": len(lines),
                        "hours": round(sum(_num(x.get("hours")) for x in lines), 2), "lastImport": last and last.get("imported")})
        return out

    def import_batch(self, pid, covers, lines, source="", layout=""):
        """Replace the lines posted in covers[from..to] with `lines` as one new batch. Returns the batch."""
        lo, hi = (covers or {}).get("from"), (covers or {}).get("to")
        if not (isinstance(lo, str) and ISO.match(lo) and isinstance(hi, str) and ISO.match(hi) and lo <= hi):
            raise CostError("The import needs the range of posting dates it covers (from ≤ to).")
        clean = [_clean_line(x, i) for i, x in enumerate(lines or [])]
        led = self.load(pid)
        bid = "b_" + secrets.token_hex(4)
        inside = lambda x: lo <= _posted(x) <= hi
        replaced = [x for x in led["lines"] if inside(x)]
        if replaced:
            trash = self.dir / ".trash"
            trash.mkdir(parents=True, exist_ok=True)
            fsutil.write_text(trash / f"{pid}-{bid}.json", json.dumps({"batch": bid, "replaced": replaced}), mode=PRIVATE)
        for x in clean:
            x["batch"] = bid
        led["lines"] = [x for x in led["lines"] if not inside(x)] + clean
        batch = {"id": bid, "source": str(source)[:200], "layout": str(layout)[:80], "imported": time.strftime("%Y-%m-%dT%H:%M:%S"),
                 "covers": {"from": lo, "to": hi}, "rows": len(clean), "replaced": len(replaced),
                 "hours": round(sum(_num(x.get("hours")) for x in clean), 2), "pay": round(sum(_num(x.get("pay")) for x in clean), 2)}
        led["batches"].append(batch)
        led["postedThrough"] = max(filter(None, [led.get("postedThrough"), hi]))
        self._save(led)
        return batch

    def undo(self, pid, bid):
        """Undo the most recent import of a project: its lines go, the lines it replaced come back."""
        led = self.load(pid)
        if not led["batches"] or led["batches"][-1]["id"] != bid:
            raise CostError("Only the most recent import can be undone.")
        tf = self.dir / ".trash" / f"{pid}-{bid}.json"
        back = (fsutil.read_json(tf, default={}) or {}).get("replaced") or []
        gone = [x for x in led["lines"] if x.get("batch") == bid]
        led["lines"] = [x for x in led["lines"] if x.get("batch") != bid] + back
        b = led["batches"].pop()
        led["postedThrough"] = max([x["covers"]["to"] for x in led["batches"] if x.get("covers")] or [None], key=lambda v: v or "")
        trash = self.dir / ".trash"
        trash.mkdir(parents=True, exist_ok=True)  # the undone lines are kept, never deleted outright
        fsutil.write_text(trash / f"{pid}-{bid}-undone.json", json.dumps({"batch": b, "lines": gone}), mode=PRIVATE)
        if tf.exists():
            fsutil.replace(tf, trash / f"{pid}-{bid}-restored.json")
        self._save(led)
        return {"removed": len(gone), "restored": len(back), "postedThrough": led["postedThrough"]}


def _num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return 0.0
    return f if f == f and abs(f) != float("inf") else 0.0


def _posted(x):
    return x.get("post") or x.get("date") or ""


TEXT_FIELDS = ("name", "emp", "payType", "category", "nwa", "task", "doc", "desc")
NUM_FIELDS = ("hours", "pay", "fringe")


def _clean_line(x, i):
    if not isinstance(x, dict):
        raise CostError(f"Line {i + 1} isn't an object.")
    d, p = x.get("date"), x.get("post") or x.get("date")
    if not (isinstance(d, str) and ISO.match(d)) or not (isinstance(p, str) and ISO.match(p)):
        raise CostError(f"Line {i + 1} needs a date (YYYY-MM-DD).")
    out = {"date": d, "post": p}
    for k in NUM_FIELDS:
        if x.get(k) not in (None, ""):
            out[k] = round(_num(x[k]), 4)
    for k in TEXT_FIELDS:
        if x.get(k):
            out[k] = str(x[k])[:200]
    if "row" in x:
        out["row"] = int(_num(x["row"]))
    return out
