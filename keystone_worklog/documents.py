"""Plain text out of the documents in the vault: PDF, Word, PowerPoint, Excel, emails, and any plain-text file
(code, Jupyter notebooks, logs, solver input/output such as Abaqus .inp/.msg/.dat/.sta, configs, CSV…).

Everything runs on this computer. Each extractor returns sections – [{"heading": "Page 3", "text": "…"}] –
so search results can say where in the file a passage is. Extracted text is cached in data/index/docs/
(keyed by the file's path and modification time), so a file is only read again after it changes.

PDF, Word, PowerPoint and Excel need optional packages (pip install "keystone-worklog[documents]");
emails, text, CSV, HTML and RTF work with the standard library alone.
"""
import csv
import email
import email.policy
import hashlib
import html
import io
import json
import os
import re
import threading
from html.parser import HTMLParser
from pathlib import Path

try:
    import pypdf
except ImportError:
    pypdf = None
try:
    import docx  # python-docx
except ImportError:
    docx = None
try:
    import pptx  # python-pptx
except ImportError:
    pptx = None
try:
    import openpyxl
except ImportError:
    openpyxl = None

# extension → (label, library needed or None)
TYPES = {
    ".pdf": ("PDF", pypdf), ".docx": ("Word", docx), ".docm": ("Word", docx), ".pptx": ("PowerPoint", pptx), ".pptm": ("PowerPoint", pptx),
    ".xlsx": ("Excel", openpyxl), ".xlsm": ("Excel", openpyxl), ".eml": ("Email", True),
    ".txt": ("Text", True), ".text": ("Text", True), ".csv": ("CSV", True), ".tsv": ("CSV", True),
    ".html": ("HTML", True), ".htm": ("HTML", True), ".rtf": ("RTF", True),
}
# plain-text files, read as text once their first bytes confirm they really are text
TEXT_TYPES = {
    ".py": "Python", ".pyw": "Python", ".ipynb": "Jupyter notebook", ".log": "Log", ".out": "Output", ".lis": "Output", ".rpt": "Report",
    ".inp": "Input deck", ".msg": "Message file", ".dat": "Data file", ".sta": "Status file", ".pes": "Input deck",  # Abaqus job files
    ".json": "JSON", ".yaml": "YAML", ".yml": "YAML", ".xml": "XML", ".toml": "Config", ".ini": "Config", ".cfg": "Config", ".conf": "Config",
    ".m": "MATLAB", ".r": "R", ".jl": "Julia", ".f": "Fortran", ".for": "Fortran", ".f90": "Fortran", ".f77": "Fortran",
    ".c": "C", ".h": "C", ".cpp": "C++", ".hpp": "C++", ".cs": "C#", ".java": "Java", ".js": "JavaScript", ".ts": "TypeScript",
    ".sh": "Shell", ".bash": "Shell", ".bat": "Batch", ".cmd": "Batch", ".ps1": "PowerShell", ".sql": "SQL", ".tex": "LaTeX", ".bib": "BibTeX",
    ".k": "LS-DYNA", ".bdf": "Nastran", ".nas": "Nastran", ".fem": "Nastran", ".cdb": "ANSYS", ".mac": "ANSYS", ".apdl": "ANSYS",
}
LOG_TYPES = {".log", ".out", ".lis", ".msg", ".dat", ".sta"}  # get an "Errors, warnings and status" section
# never sniffed: binary or geometry files that only add noise (Abaqus .odb/.prt/.sim/.res/.mdl/.stt, archives, media…)
BINARY_EXTS = set("""
.odb .prt .sim .res .mdl .stt .abq .sel .lck .023 .fil .sup .cae .jnl .rec .exe .dll .so .dylib .bin .o .a .lib .obj .pyc .pyd .pkl .pickle
.npy .npz .h5 .hdf5 .mat .db .sqlite .sqlite3 .zip .gz .tgz .bz2 .xz .7z .tar .rar .iso .dmg .msi .jar .class .woff .woff2 .ttf .otf .ico
.png .jpg .jpeg .gif .webp .bmp .tif .tiff .svg .psd .ai .eps .mp3 .wav .flac .mp4 .mov .avi .mkv .webm .doc .xls .ppt .odt .ods .odp
.dwg .stp .step .igs .iges .stl .x_t .sat .3dm .md .canvas .lnk
""".split())
LEGACY = {".doc": "Word (old format)", ".xls": "Excel (old format)", ".ppt": "PowerPoint (old format)", ".odt": "Document", ".ods": "Spreadsheet", ".odp": "Slides"}
ALL_TEXT = True                   # also any other file that looks like text (Settings → "other plain-text files")
MAX_BYTES = 60 * 1024 * 1024      # bigger PDF / Office files are listed but not read
MAX_TEXT = 3_000_000              # characters kept per file (~2,500 pages of prose)
MAX_ROWS = 5000                   # spreadsheet rows read per sheet
SMALL_TEXT = 1_500_000            # plain-text files up to this size are read whole (solver output and logs: SMALL_LOG)…
SMALL_LOG = 500_000
HEAD, TAIL = 200_000, 100_000     # …bigger ones: the start, the end, and a scan of the whole file for errors and warnings
LOG_HEAD, LOG_TAIL = 60_000, 40_000  # solver output and logs are mostly numbers: keep less of them
SCAN_LIMIT = 4 * 1024 ** 3        # stop scanning a giant output file after 4 GB
VERSION = 3                       # bump to re-extract everything after an extractor change


class DocError(Exception):
    pass


def supported():
    """Extensions that can be read with the packages installed here."""
    return {ext for ext, (_, lib) in TYPES.items() if lib}


def missing():
    """{label: pip package} for document types that need a package that isn't installed."""
    need = {"PDF": ("pypdf", pypdf), "Word": ("python-docx", docx), "PowerPoint": ("python-pptx", pptx), "Excel": ("openpyxl", openpyxl)}
    return {label: pkg for label, (pkg, lib) in need.items() if lib is None}


def kind(path):
    """A label from the extension alone (no file access)."""
    ext = Path(path).suffix.lower()
    return TYPES[ext][0] if ext in TYPES else TEXT_TYPES.get(ext) or LEGACY.get(ext) or "Text"


_sniffed = {}


def _is_text(b):
    if b.startswith((b"\xff\xfe", b"\xfe\xff", b"\xef\xbb\xbf")):
        return True
    if not b or b"\x00" in b:
        return False
    bad = sum(1 for c in b if c < 9 or 13 < c < 32 and c != 27)
    return bad / len(b) < 0.02


def looks_like_text(path, mtime=None):
    """True if the file's first 8 KB are text (remembered per path + mtime)."""
    key = (str(path), mtime)
    if key not in _sniffed:
        try:
            with open(path, "rb") as f:
                head = f.read(8192)
        except OSError:
            return False
        if len(_sniffed) > 100_000:
            _sniffed.clear()
        _sniffed[key] = _is_text(head)
    return _sniffed[key]


def doc_kind(path, mtime=None):
    """The label if Keystone can read this file (needs the real path, for the text check), else None."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext in TYPES:
        return TYPES[ext][0] if TYPES[ext][1] else None
    if ext in BINARY_EXTS or p.name.startswith(("~$", ".")):
        return None
    if ext in TEXT_TYPES or ALL_TEXT:
        return (TEXT_TYPES.get(ext) or "Text") if looks_like_text(p, mtime) else None
    return None


def is_document(path, mtime=None):
    return doc_kind(path, mtime) is not None


def listable(path, mtime=None):
    """Shown in the Notes tree: readable documents, plus old Office files and Outlook .msg (opened in their app)."""
    ext = Path(path).suffix.lower()
    return ext in LEGACY or ext == ".msg" or is_document(path, mtime)


# ---------- helpers ----------
def _clean(text):
    text = (text or "").replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _decode(data):
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


class _HTMLText(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "table", "section", "article", "blockquote", "pre"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "head"):
            self.skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n")
        elif tag in ("td", "th"):
            self.out.append(" | ")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "head"):
            self.skip = max(0, self.skip - 1)
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def html_to_text(s):
    p = _HTMLText()
    try:
        p.feed(s)
        p.close()
    except Exception:
        return _clean(re.sub(r"<[^>]+>", " ", html.unescape(s)))
    return _clean("".join(p.out))


def rtf_to_text(s):
    """A small RTF reader: drops control words and hidden groups, keeps the text."""
    out, stack, skip, i = [], [], False, 0
    hidden = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "object", "header", "footer", "themedata", "datastore", "listtable", "listoverridetable", "rsidtbl", "generator"}
    tokens = re.finditer(r"\\([a-z]{1,32})(-?\d{1,10})? ?|\\'([0-9a-f]{2})|\\([^a-z])|([{}])|[\r\n]+|([^\\{}\r\n]+)", s, re.I)
    for m in tokens:
        word, arg, hexc, sym, brace, text = m.groups()
        if brace == "{":
            stack.append(skip)
        elif brace == "}":
            skip = stack.pop() if stack else False
        elif skip:
            continue
        elif word:
            if word in hidden:
                skip = True
            elif word in ("par", "line", "row", "sect", "page"):
                out.append("\n")
            elif word in ("tab", "cell"):
                out.append("\t")
            elif word == "u" and arg:
                out.append(chr(int(arg) % 65536))
        elif hexc:
            out.append(bytes([int(hexc, 16)]).decode("cp1252", errors="replace"))
        elif sym:
            if sym == "*":
                skip = True
            elif sym in "\\{}":
                out.append(sym)
            elif sym == "~":
                out.append(" ")
        elif text:
            out.append(text)
    return _clean("".join(out))


EMAIL_HEAD = re.compile(r"^(from|sent|date|to|cc|subject)\s*:", re.I)


def _email_heading(fields):
    subj = (fields.get("subject") or "").strip() or "(no subject)"
    who = (fields.get("from") or "").strip()
    when = (fields.get("date") or fields.get("sent") or "").strip()
    return "Email: " + subj + (f" — {who}" if who else "") + (f", {when}" if when else "")


def _text_sections(text, title):
    """A text file; an email saved as text (From/Sent/To/Subject at the top) gets an email heading."""
    text = _clean(text)
    lines, fields = text.split("\n"), {}
    for line in lines[:12]:
        m = EMAIL_HEAD.match(line)
        if m:
            fields.setdefault(m.group(1).lower(), line.split(":", 1)[1].strip())
    if "subject" in fields and ("from" in fields or "to" in fields):
        return [{"heading": _email_heading(fields), "text": text}]
    return [{"heading": "", "text": text}]


# ---------- extractors (each takes bytes, so email attachments work the same way) ----------
def _pdf(data, name):
    try:
        r = pypdf.PdfReader(io.BytesIO(data))
        if r.is_encrypted and not r.decrypt(""):
            raise DocError("the PDF is password-protected")
    except DocError:
        raise
    except Exception as e:
        raise DocError(f"not a readable PDF ({e})") from None
    out = []
    for i, page in enumerate(r.pages, 1):
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        out.append({"heading": f"Page {i}", "text": _clean(t)})
    if out and not any(s["text"] for s in out):
        raise DocError("the PDF has no text layer (scanned without OCR)")
    return out


def _docx(data, name):
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise DocError(f"not a readable Word file ({e})") from None
    out, head, buf = [], "", []
    body = d.element.body
    for el in body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = docx.text.paragraph.Paragraph(el, d)
            style = (p.style.name if p.style is not None else "") or ""
            if style.lower().startswith(("heading", "title")) and p.text.strip():
                out.append({"heading": head, "text": "\n".join(buf)})
                head, buf = p.text.strip(), []
            elif p.text.strip():
                buf.append(p.text)
        elif tag == "tbl":
            t = docx.table.Table(el, d)
            for row in t.rows:
                cells, seen = [], set()
                for c in row.cells:  # merged cells repeat; keep each once
                    if id(c._tc) not in seen:
                        seen.add(id(c._tc))
                        cells.append(c.text.strip())
                if any(cells):
                    buf.append(" | ".join(cells))
            buf.append("")
    out.append({"heading": head, "text": "\n".join(buf)})
    return [{"heading": s["heading"], "text": _clean(s["text"])} for s in out if s["text"].strip() or s["heading"]]


def _pptx(data, name):
    try:
        prs = pptx.Presentation(io.BytesIO(data))
    except Exception as e:
        raise DocError(f"not a readable PowerPoint file ({e})") from None
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title, parts = "", []
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text.strip()
        except Exception:
            pass

        def walk(shapes):
            for sh in shapes:
                if getattr(sh, "shape_type", None) == 6 and hasattr(sh, "shapes"):  # group
                    walk(sh.shapes)
                elif getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                    t = sh.text_frame.text.strip()
                    if t and t != title:
                        parts.append(t)
                elif getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        if any(cells):
                            parts.append(" | ".join(cells))
        walk(slide.shapes)
        try:
            if slide.has_notes_slide:
                n = slide.notes_slide.notes_text_frame.text.strip()
                if n:
                    parts.append("Speaker notes: " + n)
        except Exception:
            pass
        out.append({"heading": f"Slide {i}" + (f": {title}" if title else ""), "text": _clean("\n".join(parts))})
    return out


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if hasattr(v, "isoformat"):
        v = v.isoformat()[:10] if getattr(v, "hour", 0) == 0 and getattr(v, "minute", 0) == 0 else v.isoformat(sep=" ", timespec="minutes")
    return str(v).strip()


def _rows_text(rows):
    """Rows as "Header: value; …" lines when the first row looks like headers, else "a | b | c"."""
    rows = [[_cell(v) for v in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    head = rows[0]
    textual = sum(1 for h in head if h and not re.fullmatch(r"[-\d.,%$ ]+", h))
    width = max(sum(1 for v in r if v) for r in rows)
    lines = []
    if len(rows) > 1 and width > 1 and textual >= max(1, len([h for h in head if h]) // 2):  # a one-column list has no header row
        lines.append(" | ".join(h for h in head if h))
        for r in rows[1:]:
            pairs = [f"{head[i] if i < len(head) and head[i] else f'Col {i + 1}'}: {v}" for i, v in enumerate(r) if v]
            lines.append("; ".join(pairs))
    else:
        lines = [" | ".join(v for v in r if v) for r in rows]
    return "\n".join(lines)


def _xlsx(data, name):
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise DocError(f"not a readable Excel file ({e})") from None
    out = []
    try:
        for ws in wb.worksheets:
            rows = []
            for n, r in enumerate(ws.iter_rows(values_only=True)):
                if n >= MAX_ROWS:
                    rows.append([f"… more rows not read (limit {MAX_ROWS})"])
                    break
                rows.append(r)
            out.append({"heading": f"Sheet: {ws.title}", "text": _clean(_rows_text(rows))})
    finally:
        wb.close()
    return out


def _csv(data, name):
    text = _decode(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if name.lower().endswith(".tsv") else csv.excel
    rows = []
    for n, r in enumerate(csv.reader(io.StringIO(text), dialect)):
        if n >= MAX_ROWS:
            rows.append([f"… more rows not read (limit {MAX_ROWS})"])
            break
        rows.append(r)
    return [{"heading": "", "text": _clean(_rows_text(rows))}]


def _eml(data, name, depth=0):
    try:
        msg = email.message_from_bytes(data, policy=email.policy.default)
    except Exception as e:
        raise DocError(f"not a readable email ({e})") from None
    fields = {k: str(msg.get(k, "") or "") for k in ("from", "to", "cc", "date", "subject")}
    head = _email_heading(fields)
    plain = htmltext = None
    attachments = []
    for part in msg.walk():
        if part.is_multipart():
            continue
        fname = part.get_filename()
        disp = part.get_content_disposition()
        ctype = part.get_content_type()
        if fname or disp == "attachment":
            attachments.append((fname or "attachment", part))
        elif ctype == "text/plain" and plain is None:
            plain = part.get_content()
        elif ctype == "text/html" and htmltext is None:
            htmltext = part.get_content()
    body = plain if plain and plain.strip() else html_to_text(htmltext or "")
    meta = "\n".join(f"{k.capitalize()}: {v}" for k, v in fields.items() if v)
    if attachments:
        meta += "\nAttachments: " + ", ".join(a for a, _ in attachments)
    out = [{"heading": head, "text": _clean(meta + "\n\n" + (body or ""))}]
    for fname, part in attachments[:20]:  # readable attachments are searched too (one level deep)
        if depth or Path(fname).suffix.lower() not in supported():
            continue
        try:
            got = extract_bytes(part.get_payload(decode=True) or b"", fname, depth + 1)
        except DocError:
            continue
        out += [{"heading": f"Attachment: {fname}" + (f" › {x['heading']}" if x["heading"] else ""), "text": x["text"]} for x in got if x["text"]]
    return out


# ---------- plain-text files ----------
ALERT = re.compile(rb"\*\*\*\s*(?:ERROR|WARNING)|\b(?:ERROR|FATAL|WARNING)\b|Traceback \(most recent call last\)|"
                   rb"HAS (?:NOT BEEN )?COMPLETED|ANALYSIS (?:COMPLETE|TERMINATED|HAS)|NOT CONVERGED|TOO MANY (?:ATTEMPTS|INCREMENTS)|"
                   rb"TIME INCREMENT REQUIRED IS LESS THAN THE MINIMUM|Abaqus/\w+ (?:exited with error|Analysis exited)")


def _clean_code(text):
    """Like _clean, but keeps indentation and column alignment (code, solver tables)."""
    text = (text or "").replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{4,}", "\n\n\n", text).strip("\n")


class _Alerts:
    """Error / warning / completion lines, each counted once (solvers repeat the same warning thousands of times)."""

    def __init__(self, cap=200):
        self.seen, self.order, self.cap = {}, [], cap

    def add(self, line):
        key = line.strip()[:300]
        if key in self.seen:
            self.seen[key] += 1
        elif len(self.order) < self.cap:
            self.seen[key] = 1
            self.order.append(key)

    def text(self):
        return "\n".join(f"{k}  (×{self.seen[k]})" if self.seen[k] > 1 else k for k in self.order)


def _python_sections(text):
    """A Python file, split at top-level def / class so results can say where."""
    out, head, buf = [], "Module code", []
    for i, line in enumerate(text.split("\n"), 1):
        m = re.match(r"(async\s+def|def|class)\s+(\w+)", line)
        if m and any(x.strip() for x in buf):
            out.append({"heading": head, "text": _clean_code("\n".join(buf))})
            buf = []
        if m:
            head = f"{m.group(1).split()[-1]} {m.group(2)} (line {i})"
        buf.append(line)
    out.append({"heading": head, "text": _clean_code("\n".join(buf))})
    return [x for x in out if x["text"]]


def _notebook(data):
    try:
        nb = json.loads(_decode(data))
        cells = nb["cells"] if "cells" in nb else [c for w in nb.get("worksheets", []) for c in w.get("cells", [])]
    except (ValueError, KeyError, TypeError) as e:
        raise DocError(f"not a readable Jupyter notebook ({e})") from None
    text = lambda v: "".join(v) if isinstance(v, list) else str(v or "")  # noqa: E731
    out = []
    for i, c in enumerate(cells, 1):
        body = text(c.get("source", c.get("input", "")))
        outs = []
        for o in c.get("outputs", []) or []:
            if o.get("output_type") == "stream":
                outs.append(text(o.get("text")))
            elif o.get("output_type") == "error":
                outs.append(f"{o.get('ename', 'Error')}: {o.get('evalue', '')}")
            elif "text/plain" in (o.get("data") or {}):
                outs.append(text(o["data"]["text/plain"]))
        if outs:
            body += "\n\nOutput:\n" + "\n".join(x[:2000] for x in outs)
        if body.strip():
            out.append({"heading": f"Cell {i} · {c.get('cell_type', 'code')}", "text": _clean_code(body)})
    return out


def _text_file(p):
    size, ext = p.stat().st_size, p.suffix.lower()
    with open(p, "rb") as f:
        head = f.read(8192)
    if not _is_text(head):
        if head.startswith(b"\xd0\xcf\x11\xe0"):
            raise DocError("this is an Outlook .msg email or an old Office file – save it as .eml / text, or .docx/.xlsx, to search it")
        raise DocError("binary file, not text")
    if ext == ".ipynb":
        if size > MAX_BYTES:
            raise DocError(f"the notebook is larger than {MAX_BYTES // (1024 * 1024)} MB")
        return _notebook(p.read_bytes())
    log = ext in LOG_TYPES
    if size <= (min(SMALL_LOG, SMALL_TEXT) if log else SMALL_TEXT):
        data = p.read_bytes()
        text = _decode(data)
        secs = _python_sections(text) if ext in (".py", ".pyw") else _text_sections(text, p.name) if ext not in TEXT_TYPES else [{"heading": "", "text": _clean_code(text)}]
        if log:
            a = _Alerts()
            for line in data.splitlines():
                if ALERT.search(line):
                    a.add(line.decode("utf-8", errors="replace"))
            if a.order:
                secs.insert(0, {"heading": "Errors, warnings and status", "text": a.text()})
        return secs
    # a big file (solver output can be gigabytes): the start, the end, and every error / warning line in between
    hd, tl = (LOG_HEAD, LOG_TAIL) if log else (HEAD, TAIL)
    with open(p, "rb") as f:
        start = f.read(hd)
        f.seek(max(hd, size - tl))
        end = f.read(tl)
    a, kw, read = _Alerts(), _Alerts(400), 0
    deck = ext in (".inp", ".pes")
    with open(p, "rb") as f:
        for line in f:
            read += len(line)
            if ALERT.search(line):
                a.add(line.decode("utf-8", errors="replace"))
            elif deck and line.startswith(b"*"):  # the model's structure: *Material, *Step, element types, comments…
                kw.add(line.decode("utf-8", errors="replace"))
            if read > SCAN_LIMIT:
                break
    mb = lambda n: f"{n / 1e6:.0f} MB"  # noqa: E731
    secs = [{"heading": "Errors, warnings and status", "text": a.text()}] if a.order else []
    if kw.order:
        secs.append({"heading": "Model keywords", "text": kw.text()})
    secs += [{"heading": "Start of file", "text": _clean_code(start.decode("utf-8", errors="replace"))},
             {"heading": f"End of file (the middle {mb(size - len(start) - len(end))} of this {mb(size)} file was only scanned for errors and warnings)",
              "text": _clean_code(end.decode("utf-8", errors="replace").split("\n", 1)[-1])}]
    return secs


def extract_bytes(data, name, depth=0):
    ext = Path(name).suffix.lower()
    label, lib = TYPES.get(ext, (None, None))
    if not label and _is_text(data[:8192]):  # a plain-text email attachment (.py, .log, .inp…)
        return _notebook(data) if ext == ".ipynb" else [{"heading": "", "text": _clean_code(_decode(data[:SMALL_TEXT]))}]
    if not label:
        raise DocError(f"{ext or 'this'} files aren't supported")
    if not lib:
        raise DocError(f"reading {label} files needs the {missing().get(label, 'optional')} package")
    if ext == ".pdf":
        out = _pdf(data, name)
    elif ext in (".docx", ".docm"):
        out = _docx(data, name)
    elif ext in (".pptx", ".pptm"):
        out = _pptx(data, name)
    elif ext in (".xlsx", ".xlsm"):
        out = _xlsx(data, name)
    elif ext in (".csv", ".tsv"):
        out = _csv(data, name)
    elif ext == ".eml":
        out = _eml(data, name, depth)
    elif ext in (".html", ".htm"):
        out = _text_sections(html_to_text(_decode(data)), name)
    elif ext == ".rtf":
        out = _text_sections(rtf_to_text(_decode(data)), name)
    else:
        out = _text_sections(_decode(data), name)
    total, kept = 0, []
    for s in out:  # cap very large files
        if total >= MAX_TEXT:
            break
        t = s["text"][: MAX_TEXT - total]
        total += len(t)
        kept.append({"heading": s["heading"], "text": t})
    return kept


def extract(path):
    p = Path(path)
    if p.suffix.lower() not in TYPES:
        return _text_file(p)
    size = p.stat().st_size
    if size > MAX_BYTES:
        raise DocError(f"the file is larger than {MAX_BYTES // (1024 * 1024)} MB")
    return extract_bytes(p.read_bytes(), p.name)


# ---------- files in the vault + the text cache ----------
def vault_documents(root):
    """{relative path: mtime_ns} for every visible file Keystone can read (notes excluded)."""
    out = {}
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if not d.startswith(".") and d not in ("__pycache__", "node_modules", ".ipynb_checkpoints")]
        for f in filenames:
            if f.startswith((".", "~$")) or f.lower().endswith(".md"):
                continue
            p = Path(dirpath) / f
            try:
                m = p.stat().st_mtime_ns
            except OSError:
                continue
            if doc_kind(p, m):
                out[p.relative_to(root).as_posix()] = m
    return out


class TextCache:
    """Extracted text per document, on disk, refreshed when the file's mtime changes."""

    def __init__(self, base_dir):
        self.dir = Path(base_dir)
        self.lock = threading.Lock()
        self._busy = set()
        self._warming = False

    def _file(self, root, rel):
        return self.dir / (hashlib.sha1(f"{root}\n{rel}".encode()).hexdigest() + ".json")

    def cached(self, root, rel, mtime):
        """The cached entry if it is current, else None. Never reads the document."""
        try:
            j = json.loads(self._file(root, rel).read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return None
        return j if j.get("mtime") == str(mtime) and j.get("v") == VERSION else None

    def get(self, root, rel, mtime=None):
        """{"sections": [...], "error": str|None}, extracting the document if needed."""
        p = Path(root) / rel
        if mtime is None:
            mtime = p.stat().st_mtime_ns
        hit = self.cached(root, rel, mtime)
        if hit:
            return hit
        try:
            sections, err = extract(p), None
        except DocError as e:
            sections, err = [], str(e)
        except Exception as e:  # a broken file must never stop indexing
            sections, err = [], f"could not read the file ({type(e).__name__}: {e})"
        entry = {"v": VERSION, "path": rel, "mtime": str(mtime), "sections": sections, "error": err}
        self.dir.mkdir(parents=True, exist_ok=True)
        f = self._file(root, rel)
        tmp = f.with_suffix(f".{threading.get_ident()}.tmp")
        tmp.write_text(json.dumps(entry), encoding="utf-8")
        os.replace(tmp, f)
        return entry

    def warm(self, root):
        """Extract every changed document in the background (for keyword search before indexing is done)."""
        with self.lock:
            if self._warming:
                return
            self._warming = True

        def run():
            try:
                for rel, m in sorted(vault_documents(root).items(), key=lambda x: -x[1]):
                    if not self.cached(root, rel, m):
                        self.get(root, rel, m)
            except Exception:
                pass
            finally:
                with self.lock:
                    self._warming = False
        threading.Thread(target=run, name="keystone-doc-text", daemon=True).start()

    def prune(self, root, keep):
        """Remove cached text of documents that are gone (or belong to a previous vault)."""
        want = {self._file(root, rel).name for rel in keep}
        for f in self.dir.glob("*.json") if self.dir.is_dir() else []:
            if f.name not in want:
                f.unlink(missing_ok=True)
