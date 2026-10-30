"""The project brief as a PowerPoint deck.

Built for a company template: every slide uses one of PowerPoint's standard layouts (Title Slide, Title and Content,
Comparison, Title Only) and puts its text in that layout's placeholders, so changing the slide master – a company
background, logo, fonts and colours – restyles every slide. No fonts or colours are fixed, except the red/amber/green of
the status, which must mean the same everywhere. Tables and charts are native PowerPoint objects (charts keep their data:
right-click → Edit Data), placed where the layout's content placeholder is.

The slides arrive already decided by the page (static/planner/report.js → deckSpec), after the user's
"Review before sharing", so this only lays them out.
"""
import base64
import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Emu, Pt

LAYOUTS = {"title": "Title Slide", "content": "Title and Content", "two": "Comparison", "titleOnly": "Title Only"}
FALLBACK = {"title": 0, "content": 1, "two": 4, "titleOnly": 5}  # positions in PowerPoint's default layout list
RAG = {"green": RGBColor(0x15, 0x80, 0x3D), "amber": RGBColor(0xB4, 0x53, 0x09), "red": RGBColor(0xC8, 0x1E, 0x1E)}
WIDE = (Emu(12192000), Emu(6858000))  # 13.33 × 7.5 in, PowerPoint's 16:9
MAX_ROWS = 14


def new_presentation(template=None):
    """A 16:9 presentation. PowerPoint's built-in template is 4:3, so its master and layouts are stretched sideways."""
    if template:
        return Presentation(template)
    prs = Presentation()
    k = WIDE[0] / prs.slide_width
    prs.slide_width, prs.slide_height = WIDE
    for shapes in [prs.slide_master.shapes] + [layout.shapes for layout in prs.slide_layouts]:
        for sh in shapes:
            # only shapes that store their own position: the rest inherit it (from the master), and reading their position
            # returns the inherited value – writing it back would pin a wrong copy onto the layout
            off = sh._element.xpath("./p:spPr/a:xfrm/a:off")
            ext = sh._element.xpath("./p:spPr/a:xfrm/a:ext")
            if off and ext:
                off[0].set("x", str(int(int(off[0].get("x")) * k)))
                ext[0].set("cx", str(int(int(ext[0].get("cx")) * k)))
    return prs


def layout(prs, kind):
    name = LAYOUTS.get(kind, "Title and Content")
    for lay in prs.slide_layouts:
        if lay.name == name:
            return lay
    return prs.slide_layouts[FALLBACK.get(kind, 1)]


def body_placeholders(slide):
    """Content placeholders other than the title, left to right then top to bottom."""
    out = [ph for ph in slide.placeholders if ph.placeholder_format.type not in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE,
                                                                                PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)]
    return sorted(out, key=lambda ph: (round((ph.left or 0) / 914400), ph.top or 0))


def fill(tf, items):
    """Bullets into a placeholder's text frame: items are str or {text, level}."""
    tf.clear()
    first = True
    for it in items:
        text, level = (it, 0) if isinstance(it, str) else (it.get("text", ""), int(it.get("level") or 0))
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.text, p.level, first = str(text), max(0, min(4, level)), False


def status_line(tf, status):
    if not status or not status.get("label"):
        return
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = f"Status: {status['label']}"
    r.font.bold = True
    if status.get("rag") in RAG:
        r.font.color.rgb = RAG[status["rag"]]


def take_area(ph):
    """The placeholder's box, then remove the empty placeholder so a table or chart can sit there instead."""
    box = (ph.left, ph.top, ph.width, ph.height)
    ph._element.getparent().remove(ph._element)
    return box


def add_table(slide, box, cols, rows):
    rows = rows[:MAX_ROWS]
    n = len(rows) + (1 if any(cols) else 0)
    left, top, width, height = box
    row_h = Emu(420000 if n <= 8 else 340000)  # rows as tall as their text, not stretched over the whole content area
    shape = slide.shapes.add_table(max(1, n), max(1, len(cols)), left, top, width, min(height, row_h * n))
    t = shape.table
    r0 = 0
    if any(cols):
        for j, c in enumerate(cols):
            t.cell(0, j).text = str(c)
        r0 = 1
    else:
        t.first_row = False
    for i, row in enumerate(rows):
        for j, v in enumerate(row[:len(cols)]):
            t.cell(i + r0, j).text = "" if v is None else str(v)
    size = Pt(14 if n <= 8 else 12)
    for r in t.rows:
        for c in r.cells:
            for p in c.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = size
    return shape


def add_chart(slide, box, chart):
    data = CategoryChartData()
    data.categories = chart.get("categories") or []
    fmt = chart.get("number") or "General"
    for s in chart.get("series") or []:
        data.add_series(s.get("name", ""), s.get("values") or [], number_format=fmt)
    kind = XL_CHART_TYPE.LINE if chart.get("type") == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, *box, data)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    if chart.get("type") == "line":
        ch.category_axis.tick_labels.font.size = Pt(9)
        for plot in ch.plots:
            for ser in plot.series:
                ser.smooth = False
    return gf


def add_image(slide, prs, png, title_shape):
    """A picture below the title, as large as fits, centred (the timeline and the summary Gantt)."""
    pic = io.BytesIO(png)
    top = (title_shape.top + title_shape.height + Emu(91440)) if title_shape is not None else Emu(1143000)
    margin = Emu(457200)
    w_max, h_max = prs.slide_width - 2 * margin, prs.slide_height - top - margin
    shape = slide.shapes.add_picture(pic, margin, top)
    k = min(w_max / shape.width, h_max / shape.height)
    shape.width, shape.height = int(shape.width * k), int(shape.height * k)
    shape.left = int((prs.slide_width - shape.width) / 2)
    return shape


def build(slides, images=None, template=None):
    """slides: report.js deckSpec(); images: {name: base64 PNG}. Returns the .pptx bytes."""
    prs = new_presentation(template)
    images = {k: base64.b64decode(v.split(",", 1)[-1]) for k, v in (images or {}).items() if v}
    for spec in slides:
        kind = spec.get("layout") or "content"
        slide = prs.slides.add_slide(layout(prs, kind))
        if slide.shapes.title is not None:
            slide.shapes.title.text = spec.get("title") or ""
        bodies = body_placeholders(slide)
        if kind == "title":
            if bodies:
                fill(bodies[0].text_frame, [line for line in str(spec.get("subtitle") or "").split("\n") if line])
                status_line(bodies[0].text_frame, spec.get("status"))
        elif kind == "content":
            ph = bodies[0] if bodies else None
            if spec.get("table") and ph is not None:
                add_table(slide, take_area(ph), spec["table"]["cols"], spec["table"]["rows"])
            elif spec.get("chart") and ph is not None:
                add_chart(slide, take_area(ph), spec["chart"])
            elif ph is not None:
                fill(ph.text_frame, spec.get("bullets") or [])
                status_line(ph.text_frame, spec.get("status"))
        elif kind == "two":
            # Comparison: [left heading, left content, right heading, right content] (Two Content has no headings)
            cols = [bodies[i:i + 2] for i in (0, 2)] if len(bodies) >= 4 else [[None, b] for b in bodies[:2]]
            for side, (head, content) in zip(("left", "right"), cols):
                part = spec.get(side) or {}
                if head is not None:
                    fill(head.text_frame, [part.get("heading", "")])
                if content is None:
                    continue
                if part.get("table"):
                    add_table(slide, take_area(content), part["table"]["cols"], part["table"]["rows"])
                else:
                    fill(content.text_frame, part.get("bullets") or [])
        elif kind == "titleOnly":
            png = images.get(spec.get("image") or "")
            if png:
                try:
                    add_image(slide, prs, png, slide.shapes.title)
                except Exception:  # an unreadable picture leaves the slide with its title (and notes) rather than failing the deck
                    pass
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(spec["notes"])
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
