"""Documents in the vault: text extraction, the text cache, indexing, search, read_document and opening files."""
import unittest
from email.message import EmailMessage
from unittest import mock

from helpers import Env, sample_data  # first: puts this checkout on sys.path
from keystone_worklog import documents as docs
from keystone_worklog import vector_index as vi
from test_vector_index import fake_vector


def make_pdf(pages):
    """A small PDF with a real text layer (what an OCR'd scan looks like to a reader)."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>",
            f"<< /Type /Pages /Kids [{' '.join(f'{4 + 2 * i} 0 R' for i in range(len(pages)))}] /Count {len(pages)} >>",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, text in enumerate(pages):
        stream = "BT /F1 12 Tf 72 720 Td 14 TL " + " ".join(f"({line}) Tj T*" for line in text.split("\n")) + " ET"
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>")
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    out, offsets = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode() + b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    return out + f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()


def make_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("Scope of work", level=1)
    d.add_paragraph("Replace the bearings on pier 3 of the Girder Bridge.")
    d.add_heading("Schedule", level=1)
    d.add_paragraph("Night closures between May and June.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Item", "Cost", "Bearings", "$84,000"
    d.save(path)


def make_pptx(path):
    import pptx
    prs = pptx.Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Drainage study"
    s1.placeholders[1].text = "Culvert C-12 is undersized for the 100-year storm."
    s1.notes_slide.notes_text_frame.text = "Mention the county's deadline."
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Next steps"
    s2.placeholders[1].text = "Survey in March"
    prs.save(path)


def make_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Task", "Hours", "Owner"])
    ws.append(["Load rating", 120, "Hiren"])
    ws.append(["Scour analysis", 64.0, "Maria"])
    wb.create_sheet("Notes").append(["Contingency held at 10 percent"])
    wb.save(path)


def make_eml(pdf_bytes):
    m = EmailMessage()
    m["From"], m["To"], m["Subject"], m["Date"] = "Maria Lopez <maria@county.gov>", "me@example.com", "Culvert sizing", "Tue, 03 Mar 2026 09:15:00 -0700"
    m.set_content("Hi,\nThe county wants the culvert memo by Friday.\nThanks")
    m.add_alternative("<p>Hi,</p><p>The county wants the <b>culvert memo</b> by Friday.</p>", subtype="html")
    m.add_attachment(pdf_bytes, maintype="application", subtype="pdf", filename="memo.pdf")
    return m.as_bytes()


@unittest.skipUnless(docs.pypdf and docs.docx and docs.pptx and docs.openpyxl, "document packages are not installed")
class ExtractTest(unittest.TestCase):
    def setUp(self):
        self.env = Env()

    def tearDown(self):
        self.env.close()

    def path(self, name):
        p = self.env.vault / name
        p.parent.mkdir(parents=True, exist_ok=True)
        return p

    def test_pdf_pages(self):
        p = self.path("spec.pdf")
        p.write_bytes(make_pdf(["Section 1 General", "Bearing replacement on pier 3\nElastomeric pads"]))
        s = docs.extract(p)
        self.assertEqual([x["heading"] for x in s], ["Page 1", "Page 2"])
        self.assertIn("Elastomeric pads", s[1]["text"])

    def test_pdf_without_text_layer_is_reported(self):
        p = self.path("scan.pdf")
        p.write_bytes(make_pdf([""]))
        with self.assertRaises(docs.DocError):
            docs.extract(p)

    def test_word_headings_and_tables(self):
        p = self.path("scope.docx")
        make_docx(p)
        s = docs.extract(p)
        self.assertEqual([x["heading"] for x in s], ["Scope of work", "Schedule"])
        self.assertIn("pier 3", s[0]["text"])
        self.assertIn("Bearings | $84,000", s[1]["text"])

    def test_powerpoint_slides_and_notes(self):
        p = self.path("deck.pptx")
        make_pptx(p)
        s = docs.extract(p)
        self.assertEqual([x["heading"] for x in s], ["Slide 1: Drainage study", "Slide 2: Next steps"])
        self.assertIn("Culvert C-12", s[0]["text"])
        self.assertIn("Speaker notes: Mention the county", s[0]["text"])

    def test_excel_rows_with_headers(self):
        p = self.path("budget.xlsx")
        make_xlsx(p)
        s = docs.extract(p)
        self.assertEqual([x["heading"] for x in s], ["Sheet: Budget", "Sheet: Notes"])
        self.assertIn("Task: Scour analysis; Hours: 64; Owner: Maria", s[0]["text"])
        self.assertIn("Contingency held at 10 percent", s[1]["text"])
        self.assertEqual(docs._rows_text([["Contingency held at 10 percent"], ["Rates per 2026 agreement"]]),
                         "Contingency held at 10 percent\nRates per 2026 agreement")  # one column: no header row

    def test_email_with_attachment(self):
        p = self.path("mail/culvert.eml")
        p.write_bytes(make_eml(make_pdf(["Culvert C-12 design memo"])))
        s = docs.extract(p)
        self.assertTrue(s[0]["heading"].startswith("Email: Culvert sizing — Maria Lopez"))
        self.assertIn("culvert memo by Friday", s[0]["text"])
        self.assertIn("Attachments: memo.pdf", s[0]["text"])
        self.assertEqual(s[1]["heading"], "Attachment: memo.pdf › Page 1")
        self.assertIn("design memo", s[1]["text"])

    def test_email_saved_as_text(self):
        p = self.path("mail/re-bearing.txt")
        p.write_text("From: Sam Ortiz\nSent: Monday, March 2, 2026 4:10 PM\nTo: Me\nSubject: RE: bearing pads\n\nThe pads ship next week.\n", encoding="utf-8")
        s = docs.extract(p)
        self.assertEqual(s[0]["heading"], "Email: RE: bearing pads — Sam Ortiz, Monday, March 2, 2026 4:10 PM")

    def test_text_csv_html_rtf(self):
        self.path("a.csv").write_text("Station,Elevation\n10+00,512.4\n", encoding="utf-8")
        self.assertIn("Station: 10+00; Elevation: 512.4", docs.extract(self.path("a.csv"))[0]["text"])
        self.path("b.html").write_text("<html><head><style>x{}</style></head><body><h1>Site</h1><p>Gate code &amp; parking</p><script>evil()</script></body></html>", encoding="utf-8")
        t = docs.extract(self.path("b.html"))[0]["text"]
        self.assertIn("Gate code & parking", t)
        self.assertNotIn("evil", t)
        self.path("c.rtf").write_text(r"{\rtf1\ansi{\fonttbl{\f0 Arial;}}\f0 Pier cap \b repair\b0\par Caf\'e9 meeting}", encoding="utf-8")
        t = docs.extract(self.path("c.rtf"))[0]["text"]
        self.assertIn("Pier cap repair", t)
        self.assertIn("Café meeting", t)
        self.assertNotIn("Arial", t)
        self.path("d.txt").write_bytes("Résumé – cp1252".encode("cp1252"))
        self.assertIn("Résumé", docs.extract(self.path("d.txt"))[0]["text"])

    def test_broken_file_is_an_error_not_a_crash(self):
        self.path("bad.docx").write_bytes(b"not a zip")
        with self.assertRaises(docs.DocError):
            docs.extract(self.path("bad.docx"))
        cache = docs.TextCache(self.env.data / "cache")
        e = cache.get(self.env.vault, "bad.docx")
        self.assertTrue(e["error"])
        self.assertEqual(cache.cached(self.env.vault, "bad.docx", (self.env.vault / "bad.docx").stat().st_mtime_ns)["error"], e["error"])

    def test_document_chunks_keep_their_location(self):
        chunks = vi.chunk_document([{"heading": "Page 7", "text": "word " * 300}], "spec.pdf")
        self.assertTrue(all(c["heading"] == "Page 7" for c in chunks))
        self.assertTrue(chunks[0]["embed"].startswith("spec.pdf › Page 7\n"))


@unittest.skipUnless(docs.pypdf and docs.docx and docs.pptx and docs.openpyxl, "document packages are not installed")
class VaultDocumentsTest(unittest.TestCase):
    def setUp(self):
        self.env = Env()
        self.env.put_data(sample_data(self.env.vault))
        self.env.write("Work/County meeting.md", "# County meeting\nThey asked about the schedule.\n")
        (self.env.vault / "Specs").mkdir()
        (self.env.vault / "Specs/bridge spec.pdf").write_bytes(make_pdf(["General requirements", "Elastomeric bearing pads per AASHTO\nPier 3 replacement"]))
        make_docx(self.env.vault / "Specs/scope.docx")
        make_pptx(self.env.vault / "drainage deck.pptx")
        make_xlsx(self.env.vault / "budget.xlsx")
        (self.env.vault / "run.sh").write_text("echo hi\n")
        self.c = self.env.client

    def tearDown(self):
        self.env.close()

    def test_tree_lists_documents(self):
        items = {i["path"]: i for i in self.c.get("/api/notes/tree").get_json()["items"]}
        self.assertEqual(items["Specs/bridge spec.pdf"]["type"], "doc")
        self.assertIn("size", items["budget.xlsx"])
        self.assertEqual(items["run.sh"]["type"], "doc")  # a script is readable text (but never launched)

    def test_keyword_search_and_read_document(self):
        r = self.env.tool("search_notes", {"query": "elastomeric bearing"})["result"]
        first = r.splitlines()[0]
        self.assertIn("[[bridge spec.pdf#page=2]]", first)
        self.assertIn("Specs/bridge spec.pdf › Page 2 (PDF; keyword)", first)
        r = self.env.tool("search_notes", {"query": "Scour analysis"})["result"]
        self.assertIn("budget.xlsx › Sheet: Budget (Excel", r)
        r = self.env.tool("read_document", {"path": "drainage deck.pptx"})["result"]
        self.assertIn("parts 1-2 of 2", r)
        self.assertIn("[Part 1 – Slide 1: Drainage study]", r)
        r = self.env.tool("read_document", {"path": "Specs/bridge spec.pdf#page=2", "start_part": 2})["result"]
        self.assertIn("Pier 3 replacement", r)
        self.assertNotIn("General requirements", r)
        self.assertIn("Documents (read with read_document): bridge spec.pdf, scope.docx", self.env.tool("list_folder", {"path": "Specs"})["result"])
        self.assertIn("echo hi", self.env.tool("read_document", {"path": "run.sh"})["result"])

    def test_documents_can_be_turned_off(self):
        self.c.put("/api/assistant/settings", json={"documents": False})
        r = self.env.tool("search_notes", {"query": "elastomeric bearing"})["result"]
        self.assertIn("Nothing in the vault matched", r)
        self.assertIn("turned off", self.env.tool("read_document", {"path": "budget.xlsx"})["error"])

    def test_files_raw_text_and_open(self):
        r = self.c.get("/api/files/raw?path=Specs/bridge%20spec.pdf")
        self.assertEqual(r.mimetype, "application/pdf")
        self.assertNotIn("attachment", r.headers.get("Content-Disposition", ""))
        r = self.c.get("/api/files/raw?path=budget.xlsx")
        self.assertIn("attachment", r.headers["Content-Disposition"])  # never rendered on Keystone's own page
        self.assertEqual(r.headers["X-Content-Type-Options"], "nosniff")
        j = self.c.get("/api/files/text?path=Specs/scope.docx").get_json()
        self.assertEqual(j["kind"], "Word")
        self.assertEqual(j["sections"][0]["heading"], "Scope of work")
        self.assertIn("attachment", self.c.get("/api/files/raw?path=run.sh").headers["Content-Disposition"])
        with mock.patch.object(self.env.app, "open_with_system") as opener:
            self.assertEqual(self.c.post("/api/files/open", json={"path": "budget.xlsx"}).status_code, 200)
            self.assertEqual(self.c.post("/api/files/open", json={"path": "Specs/scope.docx", "reveal": True}).status_code, 200)
            self.assertEqual(self.c.post("/api/files/open", json={"path": "run.sh"}).status_code, 400)  # never run programs
            self.assertEqual(self.c.post("/api/files/open", json={"path": "run.sh", "reveal": True}).status_code, 200)  # showing it is fine
            self.assertEqual(self.c.post("/api/files/open", json={"path": "../outside.pdf"}).status_code, 400)
            self.assertEqual(opener.call_count, 3)
            self.assertTrue(opener.call_args.args[1])

    @unittest.skipUnless(vi.available, "LanceDB is not installed")
    def test_documents_are_indexed_with_locations(self):
        st = self.c.get("/api/index/status?model=m1").get_json()
        self.assertEqual(st["documents"], 5)  # pdf, docx, pptx, xlsx and run.sh
        seen = {}
        for _ in range(20):
            j = self.c.post("/api/index/next", json={"model": "m1", "limit": 4}).get_json()
            if not j["files"]:
                break
            for f in j["files"]:
                seen[f["path"]] = [c["heading"] for c in f["chunks"]]
                chunks = [{**c, "vector": fake_vector(c["embed"])} for c in f["chunks"]]
                r = self.c.post("/api/index/upsert", json={"model": "m1", "path": f["path"], "mtime": f["mtime"], "chunks": chunks, "error": f.get("error")})
                self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(next(iter(seen)), "Work/County meeting.md")  # notes go first
        self.assertEqual(seen["Specs/bridge spec.pdf"], ["Page 1", "Page 2"])
        self.assertEqual(seen["drainage deck.pptx"], ["Slide 1: Drainage study", "Slide 2: Next steps"])
        self.assertEqual(self.c.get("/api/index/status?model=m1").get_json()["stale"], 0)
        r = self.env.tool("search_notes", {"query": "zzz"}, vector=fake_vector("culvert undersized storm"))["result"]
        self.assertIn("[[drainage deck.pptx]] — drainage deck.pptx › Slide 1: Drainage study (PowerPoint; meaning)", r.splitlines()[0])
        # a broken document is indexed by name and reported
        (self.env.vault / "broken.docx").write_bytes(b"nope")
        j = self.c.post("/api/index/next", json={"model": "m1"}).get_json()
        self.assertEqual(j["files"][0]["path"], "broken.docx")
        self.assertTrue(j["files"][0]["error"])
        f = j["files"][0]
        self.c.post("/api/index/upsert", json={"model": "m1", "path": f["path"], "mtime": f["mtime"], "error": f["error"],
                                               "chunks": [{**c, "vector": fake_vector(c["embed"])} for c in f["chunks"]]})
        st = self.c.get("/api/index/status?model=m1").get_json()
        self.assertEqual((st["stale"], st["errors"]), (0, 1))
        # turning documents off drops them from the index
        self.c.put("/api/assistant/settings", json={"documents": False})
        st = self.c.get("/api/index/status?model=m1").get_json()
        self.assertEqual((st["total"], st["documents"]), (st["notes"], 0))


class OfflineTest(unittest.TestCase):
    def test_offline_never_downloads_the_model(self):
        env = Env()
        try:
            env.app.OFFLINE = True
            env.app.MODEL_BUNDLE = env.tmp / "no-models"
            with mock.patch("urllib.request.urlopen") as net:
                r = env.client.get("/hf/" + vi.EMBED_FILES["hf"][0])
                self.assertEqual(r.status_code, 404)
                net.assert_not_called()
            self.assertTrue(env.client.get("/api/assistant/settings").get_json()["offline"])
        finally:
            env.close()


if __name__ == "__main__":
    unittest.main()
