"""The project brief: the weekly log, saving shared copies in the vault, and the PowerPoint deck built for a company master."""
import base64
import io
import json
import unittest

from helpers import Env  # first: puts this checkout on sys.path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from keystone_worklog import deck
except ImportError:  # python-pptx is in the [documents] extra
    deck = None

def _png():
    try:
        from PIL import Image
    except ImportError:
        return ""
    buf = io.BytesIO()
    Image.new("RGB", (900, 300), (79, 70, 229)).save(buf, "PNG")
    return base64.b64encode(buf.getvalue()).decode()


PNG = _png()

SLIDES = [
    {"layout": "title", "title": "Girder Bridge rehab", "subtitle": "Rehabilitate the bridge on CR 12\nCounty · Oct 29, 2026", "status": {"label": "At risk", "rag": "amber"}, "notes": "FE model late"},
    {"layout": "content", "title": "At a glance", "bullets": [{"text": "Behind schedule by 3 working days."}, {"text": "Finish: Dec 22", "level": 1}], "notes": "Behind schedule."},
    {"layout": "two", "title": "About the project", "left": {"heading": "Objectives", "bullets": ["Replace bearings", "Load rating"]}, "right": {"heading": "Scope", "bullets": ["Bearings", "Out of scope: painting"]}},
    {"layout": "two", "title": "Funding summary", "left": {"heading": "Where we are", "bullets": ["$58k of $156k spent."]}, "right": {"table": {"cols": ["", ""], "rows": [["Budget", "$156,000"], ["Spent", "$58,000"]]}}},
    {"layout": "titleOnly", "title": "Schedule", "image": "gantt", "notes": "Forecast Dec 22"},
    {"layout": "content", "title": "Milestones", "table": {"cols": ["Milestone", "Baseline", "Forecast"], "rows": [["60% plans", "Nov 12", "Nov 16"], ["PS&E", "Dec 17", "Dec 22"]]}},
    {"layout": "content", "title": "Funding", "chart": {"type": "line", "categories": ["Oct 1", "Oct 15", "Nov 1"], "series": [{"name": "Planned", "values": [0, 20000, 40000]}, {"name": "Actual", "values": [0, 22000, None]}], "number": "$#,##0"}},
    {"layout": "content", "title": "Cost by phase", "chart": {"type": "bar", "categories": ["Inspection", "Design"], "series": [{"name": "Baseline", "values": [15000, 70000]}, {"name": "Actual", "values": [17000, 3000]}]}},
]


@unittest.skipIf(deck is None, "python-pptx is not installed")
class DeckTest(unittest.TestCase):
    def setUp(self):
        self.prs = Presentation(io.BytesIO(deck.build(SLIDES, {"gantt": "data:image/png;base64," + PNG})))

    def test_standard_layouts_widescreen_and_notes(self):
        self.assertEqual((self.prs.slide_width, self.prs.slide_height), deck.WIDE)
        self.assertEqual([s.slide_layout.name for s in self.prs.slides],
                         ["Title Slide", "Title and Content", "Comparison", "Comparison", "Title Only", "Title and Content", "Title and Content", "Title and Content"])
        self.assertEqual([s.shapes.title.text for s in self.prs.slides], [s["title"] for s in SLIDES])
        self.assertEqual(self.prs.slides[0].notes_slide.notes_text_frame.text, "FE model late")
        # the layouts were stretched to 16:9 too, so text sits across the slide, not in the left three quarters
        title = self.prs.slides[1].shapes.title
        self.assertGreater(title.left + title.width, self.prs.slide_width * 0.85)

    def test_text_is_in_placeholders_with_no_fixed_fonts(self):
        """What a company master needs: every piece of text in a layout placeholder, so the master restyles it."""
        for slide in self.prs.slides:
            for sh in slide.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    self.assertTrue(sh.is_placeholder, f"loose text box on '{slide.shapes.title.text}': {sh.text_frame.text[:40]}")
                if sh.has_text_frame:
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            self.assertIsNone(r.font.name, "a font is fixed")
                            if r.font.color and r.font.color.type is not None:
                                self.assertTrue(r.text.startswith("Status:"), "only the status has a fixed colour")

    def test_native_tables_charts_and_pictures(self):
        kinds = [[sh.shape_type for sh in s.shapes if not sh.is_placeholder] for s in self.prs.slides]
        self.assertIn(MSO_SHAPE_TYPE.TABLE, kinds[3])
        self.assertEqual(kinds[4], [MSO_SHAPE_TYPE.PICTURE])
        self.assertIn(MSO_SHAPE_TYPE.TABLE, kinds[5])
        chart = next(sh for sh in self.prs.slides[6].shapes if sh.has_chart).chart
        self.assertEqual([s.name for s in chart.plots[0].series], ["Planned", "Actual"])
        self.assertEqual(list(chart.plots[0].categories), ["Oct 1", "Oct 15", "Nov 1"])
        table = next(sh for sh in self.prs.slides[5].shapes if sh.has_table).table
        self.assertEqual(table.cell(1, 0).text, "60% plans")
        # a table or chart takes the place of the content placeholder – no empty "Click to add text" left behind
        self.assertFalse([ph for ph in self.prs.slides[5].placeholders if ph.placeholder_format.idx == 1])
        status = self.prs.slides[0].placeholders[1].text_frame.paragraphs[-1]
        self.assertEqual(status.runs[0].text, "Status: At risk")

    def test_a_company_master_restyles_it(self):
        """Build on another template (as if it were the company's): its layouts are used by name and everything still fits."""
        tpl = Presentation()
        tpl.slide_width, tpl.slide_height = deck.WIDE
        buf = io.BytesIO()
        tpl.save(buf)
        buf.seek(0)
        prs = Presentation(io.BytesIO(deck.build(SLIDES, {"gantt": PNG}, template=buf)))
        self.assertEqual(len(prs.slides), len(SLIDES))
        for slide in prs.slides:
            for sh in slide.shapes:
                if sh.left is not None:
                    self.assertLessEqual(sh.left + sh.width, prs.slide_width + 1, f"off the slide: {slide.shapes.title.text}")


class BriefAPITest(unittest.TestCase):
    def setUp(self):
        self.env = Env()
        self.env.put_data({"settings": {"vaultPath": str(self.env.vault), "vaultFolder": "Keystone"}, "projects": []})
        self.c = self.env.client
        plan = {"name": "P", "tasks": [{"uid": 1, "name": "A"}], "resources": [], "assignments": []}
        self.pid = self.c.post("/api/plans", json={"plan": plan}).get_json()["id"]

    def tearDown(self):
        self.env.close()

    def test_the_weekly_log(self):
        self.assertEqual(self.c.get(f"/api/plans/{self.pid}/log").get_json()["entries"], [])
        entries = [{"id": "l_1", "at": "2026-10-28T10:00", "kind": "win", "text": "County approved", "internal": False, "junk": 1},
                   {"id": "l_2", "at": "2026-10-28T11:00", "kind": "problem", "text": "Pads late", "internal": True, "taskUids": [1]}]
        self.assertEqual(self.c.put(f"/api/plans/{self.pid}/log", json={"entries": entries}).get_json()["count"], 2)
        got = self.c.get(f"/api/plans/{self.pid}/log").get_json()["entries"]
        self.assertEqual([e["text"] for e in got], ["County approved", "Pads late"])
        self.assertNotIn("junk", got[0])
        self.assertTrue(got[1]["internal"])
        bad = self.c.put(f"/api/plans/{self.pid}/log", json={"entries": [{"id": "x", "at": "now", "kind": "win"}]})
        self.assertEqual(bad.status_code, 400)
        self.assertEqual(self.c.put(f"/api/plans/{self.pid}/log", json={"entries": [{"id": "l_3", "at": "2026-10-28T10:00", "kind": "gossip"}]}).status_code, 400)
        # the log goes to .trash with its plan, never deleted outright; and it isn't listed as a plan
        self.assertEqual(len(self.c.get("/api/plans").get_json()["plans"]), 1)
        self.c.post(f"/api/plans/{self.pid}/delete", json={})
        self.assertTrue(any(p.name.endswith(".log.json") for p in (self.env.data / "plans" / ".trash").iterdir()))

    def test_saving_a_shared_copy_in_the_vault(self):
        data = base64.b64encode(b"<!doctype html><p>brief</p>").decode()
        path = "Keystone/Projects/P/Shared/2026-10-29 P brief.html"
        r1 = self.c.post("/api/reports/save", json={"path": path, "data": data}).get_json()
        r2 = self.c.post("/api/reports/save", json={"path": path, "data": data}).get_json()
        self.assertEqual(r1["path"], path)
        self.assertEqual(r2["path"], "Keystone/Projects/P/Shared/2026-10-29 P brief (2).html")  # never overwrites
        self.assertEqual((self.env.vault / path).read_bytes(), b"<!doctype html><p>brief</p>")
        for bad in ({"path": path.replace(".html", ".exe"), "data": data}, {"path": "../outside.html", "data": data}, {"path": ".obsidian/x.html", "data": data}, {"path": path, "data": "!!!"}):
            self.assertEqual(self.c.post("/api/reports/save", json=bad).status_code, 400, bad)

    @unittest.skipIf(deck is None, "python-pptx is not installed")
    def test_powerpoint_endpoint(self):
        r = self.c.post("/api/reports/pptx", json={"slides": SLIDES, "images": {"gantt": PNG}, "name": "P brief: Oct"})
        self.assertEqual(r.status_code, 200)
        self.assertIn('filename="P brief_ Oct.pptx"', r.headers["Content-Disposition"])
        self.assertEqual(len(Presentation(io.BytesIO(r.data)).slides), len(SLIDES))
        self.assertEqual(self.c.post("/api/reports/pptx", json={"slides": []}).status_code, 400)


if __name__ == "__main__":
    unittest.main()
